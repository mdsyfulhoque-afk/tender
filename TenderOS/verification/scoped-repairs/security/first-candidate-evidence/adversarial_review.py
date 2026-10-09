"""Independent synthetic security checks; no external service/model calls.

Source is imported only from the immutable candidate path supplied by parent.
All writes and synthetic SQLite/uploads occur within this harness directory.
"""
import argparse
import hashlib
import importlib
import io
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
from datetime import date, datetime, timezone, timedelta
from zipfile import ZipFile
import xml.etree.ElementTree as ET

HERE = Path(__file__).resolve().parent
TRACE = []
CASES = []
MAIN = None


def utc():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


class Fixture:
    def __init__(self):
        from fastapi.testclient import TestClient
        self.folder = Path(tempfile.mkdtemp(prefix='synthetic-', dir=HERE / 'runtime'))
        self.db = self.folder / 'fixture.sqlite3'
        self.storage = self.folder / 'uploads'
        self.storage.mkdir()
        MAIN.DB = self.db
        MAIN.STORAGE = self.storage
        with sqlite3.connect(self.db) as db:
            db.executescript(MAIN.SCHEMA)
        self.client = TestClient(MAIN.app)

    def request(self, method, path, expected=200, **kwargs):
        result = self.client.request(method, path, **kwargs)
        TRACE.append({'at_utc': utc(), 'case': CASES[-1]['name'], 'method': method,
                      'path': path, 'status': result.status_code})
        assert result.status_code == expected, (method, path, expected, result.status_code, result.text[:300])
        return result

    def sql(self, query, args=()):
        with sqlite3.connect(self.db) as db:
            db.row_factory = sqlite3.Row
            return [dict(row) for row in db.execute(query, args)]

    def history(self):
        rows = self.sql('SELECT * FROM decisions ORDER BY id')
        return hashlib.sha256(json.dumps(rows, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

    def detail(self):
        return self.request('GET', f'/api/tenders/{self.tender}').json()

    def attest(self):
        return self.request('POST', f'/api/tenders/{self.tender}/attest-source-scope', json={
            'reviewer': 'Synthetic Independent Reviewer', 'checked_full_document_set': True,
            'note': 'Synthetic full document inventory examined and source clauses checked.'})

    def decide(self, choice='BID', expected=201):
        return self.request('POST', f'/api/tenders/{self.tender}/decisions', expected=expected, json={
            'decision': choice, 'reviewer': 'Synthetic Independent Reviewer',
            'rationale': 'Synthetic evidence and provenance examined for this independent local check.'})

    def empty_source(self):
        from reportlab.pdfgen.canvas import Canvas
        pdf = io.BytesIO()
        c = Canvas(pdf)
        c.drawString(40, 700, 'Synthetic administrative cover page; schedule version beta.')
        c.showPage()
        c.save()
        response = self.request('POST', f'/api/tenders/{self.tender}/upload-pdf', expected=201,
            files={'file': ('synthetic-addendum.pdf', pdf.getvalue(), 'application/pdf')})
        assert response.json()['candidates'] == 0
        return response.json()['source_id']

    def ready(self, initial_sources=0):
        self.org = self.request('POST', '/api/organizations', expected=201,
            json={'name': 'Synthetic Security Review Ltd'}).json()['id']
        self.tender = self.request('POST', '/api/tenders', expected=201, json={
            'organization_id': self.org, 'title': 'Synthetic bounded security review'}).json()['id']
        self.sources = [self.empty_source() for _ in range(initial_sources)]
        clause = 'Three signed synthetic certificates must be held.'
        self.requirement = self.request('POST', f'/api/tenders/{self.tender}/requirements', expected=201,
            json={'text': clause, 'source_page': 2, 'source_quote': clause}).json()['id']
        self.request('PUT', f'/api/requirements/{self.requirement}/review', json={
            'text': clause, 'mandatory': True, 'source_page': 2, 'source_quote': clause})
        self.evidence = self.request('POST', '/api/evidence', expected=201, json={
            'organization_id': self.org, 'label': 'Synthetic signed certificate',
            'reference': 'SYNTHETIC-REF-1', 'expires_on': (date.today()+timedelta(days=30)).isoformat(),
            'verified': True, 'verification_note': 'Synthetic signed original checked.'}).json()['id']
        self.request('PUT', f'/api/requirements/{self.requirement}/status', json={
            'status': 'VERIFIED', 'evidence_id': self.evidence, 'notes': 'Original synthetic rationale.'})
        self.attest()
        self.decide()
        snap = self.detail()
        assert snap['compliance'] == 'READY_FOR_HUMAN_DECISION'
        assert snap['current_decision'] == 'BID'
        return self

    def assert_stale(self, old_history):
        snap = self.detail()
        assert snap['current_decision'] is None, snap
        assert all(not d['is_current'] for d in snap['decisions']), snap['decisions']
        assert self.history() == old_history, 'History changed during read/check'
        return snap


def record(name, fn):
    case = {'name': name, 'started_at_utc': utc()}
    CASES.append(case)
    try:
        fn()
        case['passed'] = True
    except Exception as exc:
        case['passed'] = False
        case['exception'] = f'{type(exc).__name__}: {exc}'[:6000]
    case['ended_at_utc'] = utc()


def zero_source_addendum():
    f = Fixture().ready()
    history = f.history()
    f.empty_source()
    snap = f.assert_stale(history)
    assert snap['compliance'] == 'UNRESOLVED' and not snap['source_scope_verified']
    f.decide(expected=409)
    assert f.history() == history
    f.attest()
    f.assert_stale(history)
    f.decide()
    snap = f.detail()
    assert snap['current_decision'] == 'BID'
    assert [d['is_current'] for d in snap['decisions']] == [True, False]


def linked_mutation(field, value):
    f = Fixture().ready()
    history = f.history()
    f.sql(f'UPDATE evidence SET {field}=? WHERE id=?', (value, f.evidence))
    f.assert_stale(history)


def requirement_notes():
    f = Fixture().ready()
    history = f.history()
    f.request('PUT', f'/api/requirements/{f.requirement}/status', json={
        'status': 'VERIFIED', 'evidence_id': f.evidence, 'notes': 'Different synthetic verification rationale.'})
    f.assert_stale(history)


def requirement_source_mapping():
    f = Fixture().ready(initial_sources=2)
    f.sql('UPDATE requirements SET source_id=? WHERE id=?', (f.sources[0], f.requirement))
    f.attest()
    f.decide()
    history = f.history()
    f.sql('UPDATE requirements SET source_id=? WHERE id=?', (f.sources[1], f.requirement))
    f.assert_stale(history)


def source_hash_mutation():
    f = Fixture().ready(initial_sources=1)
    history = f.history()
    f.sql('UPDATE sources SET sha256=? WHERE id=?', ('1'*64, f.sources[0]))
    f.assert_stale(history)


def lost_attestation():
    f = Fixture().ready()
    history = f.history()
    f.sql('UPDATE tenders SET source_scope_fingerprint=NULL WHERE id=?', (f.tender,))
    snap = f.assert_stale(history)
    assert snap['compliance'] == 'UNRESOLVED'


def reattest_same_second():
    original_now = MAIN.now
    MAIN.now = lambda: '2026-10-09T16:30:00+00:00'
    try:
        f = Fixture().ready()
        before = f.detail()
        history = f.history()
        f.attest()
        after = f.assert_stale(history)
        assert before['source_scope_fingerprint'] == after['source_scope_fingerprint']
        assert before['tender']['source_scope_at'] == after['tender']['source_scope_at']
        assert after['compliance'] == 'READY_FOR_HUMAN_DECISION'
    finally:
        MAIN.now = original_now


def unrelated_activity():
    f = Fixture().ready()
    history = f.history()
    f.request('POST', '/api/evidence', expected=201, json={
        'organization_id': f.org, 'label': 'Synthetic unlinked item', 'reference': 'UNLINKED'})
    f.request('POST', f'/api/tenders/{f.tender}/metrics', json={
        'paid_bdt': 0, 'actual_minutes': 1, 'notes': 'Synthetic metric update.'})
    other = f.request('POST', '/api/tenders', expected=201, json={
        'organization_id': f.org, 'title': 'Synthetic unrelated assessment'}).json()['id']
    f.request('POST', f'/api/tenders/{other}/requirements', expected=201,
        json={'text': 'Synthetic unrelated mandatory clause.'})
    snap = f.detail()
    assert snap['current_decision'] == 'BID' and snap['decisions'][0]['is_current']
    assert f.history() == history


def malformed_snapshot(value):
    f = Fixture().ready()
    f.sql('UPDATE decisions SET decision_snapshot=?', (value,))
    history = f.history()
    f.assert_stale(history)
    for ext in ('csv', 'xlsx', 'docx'):
        f.request('GET', f'/api/tenders/{f.tender}/export.{ext}')
    assert f.history() == history


def missing_metadata(key):
    f = Fixture().ready()
    raw = json.loads(f.sql('SELECT decision_snapshot FROM decisions')[0]['decision_snapshot'])
    assert key in raw, f'Expected existing snapshot key absent: {key}'
    raw.pop(key)
    f.sql('UPDATE decisions SET decision_snapshot=?', (json.dumps(raw),))
    history = f.history()
    f.assert_stale(history)


def superseded_latest():
    f = Fixture().ready()
    f.decide('HOLD')
    snap = f.detail()
    assert [d['is_current'] for d in snap['decisions']] == [True, False]
    f.sql('UPDATE decisions SET decision_snapshot=? WHERE id=?', ('{}', snap['decisions'][0]['id']))
    history = f.history()
    f.assert_stale(history)


def xlsx_literal_payload(payload):
    from openpyxl import load_workbook
    f = Fixture().ready()
    f.sql('UPDATE tenders SET title=? WHERE id=?', (payload, f.tender))
    f.sql('UPDATE requirements SET text=?,source_quote=?,notes=? WHERE id=?',
          (payload, payload, payload, f.requirement))
    f.sql('UPDATE evidence SET label=? WHERE id=?', (payload, f.evidence))
    f.sql('UPDATE decisions SET reviewer=?,rationale=?,created_at=?', (payload, payload, payload))
    history = f.history()
    response = f.request('GET', f'/api/tenders/{f.tender}/export.xlsx')
    book = load_workbook(io.BytesIO(response.content))
    cells = [('Compliance Matrix', 'B2'), ('Compliance Matrix', 'E8'),
             ('Compliance Matrix', 'G8'), ('Compliance Matrix', 'H8'), ('Compliance Matrix', 'I8'),
             ('Decision History', 'B2'), ('Decision History', 'C2'), ('Decision History', 'D2')]
    for sheet, address in cells:
        cell = book[sheet][address]
        assert cell.data_type == 's', (sheet, address, cell.data_type, cell.value)
        expected = "'"+payload if payload.lstrip().startswith(('=','+','-','@')) else payload
        assert cell.value == expected, (sheet, address, cell.value, expected)
    assert book['Compliance Matrix']['A8'].data_type == 'n'
    assert book['Compliance Matrix']['F8'].data_type == 'n'
    with ZipFile(io.BytesIO(response.content)) as archive:
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                assert not ET.fromstring(archive.read(name)).findall('.//m:f', ns), name
        assert not any(name.startswith('xl/externalLinks/') for name in archive.namelist())
    assert f.history() == history


def snapshot_literal_payload():
    from openpyxl import load_workbook
    f = Fixture().ready()
    payload = '=1+1 harmless snapshot'
    f.sql('UPDATE decisions SET decision_snapshot=?', (payload,))
    history = f.history()
    response = f.request('GET', f'/api/tenders/{f.tender}/export.xlsx')
    cell = load_workbook(io.BytesIO(response.content))['Decision History']['E2']
    assert cell.data_type == 's' and cell.value in (payload, "'"+payload)
    assert f.history() == history


def run():
    global MAIN
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--commit', required=True)
    parser.add_argument('--additional-metadata', nargs='*', default=[])
    args = parser.parse_args()
    source = args.source.resolve()
    actual = subprocess.check_output(['git', '-C', str(source), 'rev-parse', 'HEAD'], text=True).strip()
    assert actual == args.commit, (actual, args.commit)
    (HERE / 'runtime').mkdir(exist_ok=True)
    os.environ['TENDEROS_DATA_DIR'] = str(HERE / 'runtime' / 'import-data')
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(source / 'TenderOS'))
    MAIN = importlib.import_module('app.main')
    started = utc()
    record('zero_candidate_addendum_reattest_new_approval', zero_source_addendum)
    for field, value in [('label', 'Synthetic changed label'), ('reference', 'SYNTHETIC-REF-2'),
                         ('verification_note', 'Synthetic different signed-original verification.'),
                         ('expires_on', (date.today()+timedelta(days=60)).isoformat()),
                         ('verified', 0), ('expires_on', (date.today()-timedelta(days=1)).isoformat())]:
        record(f'linked_evidence_{field}_{value}', lambda field=field,value=value: linked_mutation(field,value))
    record('requirement_rationale_change', requirement_notes)
    record('requirement_source_id_remapping', requirement_source_mapping)
    record('source_content_hash_change', source_hash_mutation)
    record('source_attestation_missing', lost_attestation)
    record('source_reattest_same_scope_same_second', reattest_same_second)
    record('unlinked_evidence_metrics_other_tender_no_invalidation', unrelated_activity)
    for raw in ('{}', 'null', '[]', '[1]', '"synthetic string"', '1', '', '{invalid json'):
        record(f'legacy_or_malformed_snapshot_{repr(raw)}', lambda raw=raw: malformed_snapshot(raw))
    for key in ['fingerprint', 'scope_hash'] + args.additional_metadata:
        record(f'missing_snapshot_metadata_{key}', lambda key=key: missing_metadata(key))
    record('older_bid_never_resurfaces', superseded_latest)
    for payload in ('=1+1 benign', '+1+1 benign', '-1+1 benign', '@SUM(1,1) benign',
                    ' \t=1+1 benign', '\r\n=1+1 benign', '\t+1+1 benign', '\u00a0=1+1 benign',
                    'বাংলা synthetic plain text', "'ordinary literal text"):
        record(f'xlsx_dynamic_cells_literal_{repr(payload)}', lambda payload=payload: xlsx_literal_payload(payload))
    record('xlsx_history_snapshot_literal', snapshot_literal_payload)
    report = {'task_id': '/root/tenderos_patch_security', 'runtime_run_id': None,
              'runtime_run_id_exposed': False, 'strict_historical_receipt_gate_passed': False,
              'source_commit': actual, 'started_at_utc': started, 'ended_at_utc': utc(),
              'data': 'synthetic only', 'external_model_calls': False,
              'formula_evaluated': False, 'case_count': len(CASES),
              'passed': sum(c['passed'] for c in CASES), 'failed': sum(not c['passed'] for c in CASES),
              'cases': CASES,
              'source_sha256': hashlib.sha256((source/'TenderOS/app/main.py').read_bytes()).hexdigest(),
              'harness_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    (HERE/'security-results.json').write_text(json.dumps(report, indent=2, ensure_ascii=False)+'\n')
    (HERE/'synthetic-request-trace.json').write_text(json.dumps(TRACE, indent=2)+'\n')
    print(json.dumps({'case_count': report['case_count'], 'passed': report['passed'],
                      'failed': report['failed'], 'report': str(HERE/'security-results.json')}))
    return 1 if report['failed'] else 0


if __name__ == '__main__':
    raise SystemExit(run())
