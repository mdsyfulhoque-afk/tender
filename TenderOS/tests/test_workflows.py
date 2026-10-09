import io
import json
import sqlite3
from datetime import date, timedelta
import pytest
from fastapi.testclient import TestClient
from openpyxl.cell.cell import ERROR_CODES
from app import main

@pytest.fixture
def client(tmp_path,monkeypatch):
    db=tmp_path/'test.db';storage=tmp_path/'uploads';storage.mkdir()
    monkeypatch.setattr(main,'DB',db);monkeypatch.setattr(main,'STORAGE',storage)
    with sqlite3.connect(db) as cx:
        cx.executescript(main.SCHEMA)
    return TestClient(main.app)

def add_org(client,name='Test Consulting'):
    x=client.post('/api/organizations',json={'name':name});assert x.status_code==201,x.text
    return x.json()['id']

def add_tender(client,org):
    x=client.post('/api/tenders',json={'organization_id':org,'title':'Synthetic consulting assignment'})
    assert x.status_code==201,x.text
    return x.json()['id']

def add_req(client,tender,mandatory=True):
    x=client.post(f'/api/tenders/{tender}/requirements',json={'text':'Three similar projects must be documented', 'source_page':2,'source_quote':'Three similar projects must be documented'})
    assert x.status_code==201,x.text
    req=x.json()['id']
    if mandatory is not None:
        x=client.put(f'/api/requirements/{req}/review',json={'text':'Three similar projects must be documented','mandatory':mandatory,'source_page':2,'source_quote':'Three similar projects must be documented'})
        assert x.status_code==200,x.text
    return req

def evidence(client,org,verified=True,expires_on=None):
    result=client.post('/api/evidence',json={'organization_id':org,'label':'Signed certificate','reference':'CERT-001',
       'verified':verified,'verification_note':'Checked signed original' if verified else '', 'expires_on':expires_on})
    assert result.status_code==201,result.text
    return result.json()['id']

def attest(client,tender):
    result=client.post(f'/api/tenders/{tender}/attest-source-scope',json={'reviewer':'Human Reviewer','note':'I inspected the complete source, schedules and amendments','checked_full_document_set':True})
    assert result.status_code==200,result.text

def verdict(client,tender,decision='BID'):
    return client.post(f'/api/tenders/{tender}/decisions',json={'decision':decision,'reviewer':'Human Reviewer','rationale':'Checked original documents and recorded commercial choice'})

def test_mandatory_blocker_and_human_gate(client):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t)
    blocked=client.put(f'/api/requirements/{r}/status',json={'status':'NOT_HELD','notes':'No evidence exists'})
    assert blocked.status_code==200
    d=client.get(f'/api/tenders/{t}').json()
    assert d['compliance']=='BLOCKED' and len(d['blockers'])==1
    assert verdict(client,t).status_code==409
    hold=verdict(client,t,'HOLD'); assert hold.status_code==201
    ev=evidence(client,org)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev,'notes':'Checked'}).status_code==200
    attest(client,t)
    d=client.get(f'/api/tenders/{t}').json()
    assert d['compliance']=='READY_FOR_HUMAN_DECISION' and d['coverage_percent']==100
    assert verdict(client,t).status_code==201
    assert len(client.get(f'/api/tenders/{t}').json()['decisions'])==2

def test_fail_closed_when_unreviewed(client):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t,mandatory=None)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED'}).status_code==422
    assert client.get(f'/api/tenders/{t}').json()['compliance']=='UNRESOLVED'
    assert verdict(client,t).status_code==409

def test_no_verified_without_current_evidence(client):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','notes':'none'}).status_code==422
    expired=(date.today()-timedelta(days=1)).isoformat()
    ev=evidence(client,org,expires_on=expired)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev}).status_code==422
    assert client.post('/api/evidence',json={'organization_id':org,'label':'Unsigned evidence','reference':'document','verified':True}).status_code==422

def test_cross_organization_evidence_rejected(client):
    org1=add_org(client,'Org Alpha');org2=add_org(client,'Org Beta');t=add_tender(client,org1);r=add_req(client,t)
    ev=evidence(client,org2)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev}).status_code==422

def test_exports_and_metrics(client):
    org=add_org(client);t=add_tender(client,org);add_req(client,t)
    for ext,content_type in [('csv','text/csv'),('xlsx','application/vnd.openxmlformats'),('docx','application/vnd.openxmlformats')]:
        x=client.get(f'/api/tenders/{t}/export.{ext}')
        assert x.status_code==200,x.text[:200]
        assert content_type in x.headers['content-type']
        assert len(x.content)>70
    m=client.post(f'/api/tenders/{t}/metrics',json={'quoted_bdt':20000,'paid_bdt':0,'actual_minutes':90,'decision_changed':False})
    assert m.status_code==200
    assert client.get(f'/api/tenders/{t}').json()['metrics']['paid_bdt']==0

def test_pdf_has_source_provenance(client):
    reportlab=pytest.importorskip('reportlab')
    from reportlab.pdfgen.canvas import Canvas
    org=add_org(client);t=add_tender(client,org)
    buf=io.BytesIO();c=Canvas(buf);c.drawString(50,750,'Bidders must show at least three similar projects completed.');c.showPage();c.save()
    pdf=buf.getvalue()
    up=client.post(f'/api/tenders/{t}/upload-pdf',files={'file':('requirements.pdf',pdf,'application/pdf')})
    assert up.status_code==201,up.text
    assert up.json()['candidates']>=1
    d=client.get(f'/api/tenders/{t}').json()
    assert d['requirements'][0]['source_page']==1
    assert d['requirements'][0]['reviewed']==0
    assert d['requirements'][0]['mandatory'] is None
    assert len(d['sources'][0]['sha256'])==64
    assert d['compliance']=='UNRESOLVED'

def test_health_and_demo(client):
    h=client.get('/api/health');assert h.status_code==200
    assert h.json()['external_model_calls'] is False
    d=client.post('/api/demo');assert d.status_code==201
    snap=client.get(f"/api/tenders/{d.json()['tender_id']}").json()
    assert snap['compliance']=='BLOCKED'
    assert verdict(client,snap['tender']['id']).status_code==409

def test_missing_provenance_keeps_mandatory_unresolved(client):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t)
    ev=evidence(client,org)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev}).status_code==200
    # Removing the exact source quote must stop BID even when evidence has been verified.
    assert client.put(f'/api/requirements/{r}/review',json={'text':'Three similar projects must be documented','mandatory':True,'source_page':2,'source_quote':''}).status_code==200
    snap=client.get(f'/api/tenders/{t}').json()
    assert snap['compliance']=='UNRESOLVED'
    assert verdict(client,t).status_code==409

def test_changed_requirements_invalidate_prior_bid(client):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t)
    ev=evidence(client,org)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev}).status_code==200
    assert verdict(client,t).status_code==409 # missing source-inventory attestation
    attest(client,t)
    assert verdict(client,t).status_code==201
    assert client.get(f'/api/tenders/{t}').json()['current_decision']=='BID'
    assert client.post(f'/api/tenders/{t}/requirements',json={'text':'New mandatory condition from addendum', 'source_page':3,'source_quote':'New mandatory condition from addendum'}).status_code==201
    snap=client.get(f'/api/tenders/{t}').json()
    assert snap['current_decision'] is None
    assert snap['decisions'][0]['is_current'] is False
    assert verdict(client,t).status_code==409

def test_excel_formula_injection_blocked(client):
    from openpyxl import load_workbook
    org=add_org(client);t=add_tender(client,org)
    x=client.post(f'/api/tenders/{t}/requirements',json={'text':'=SUM(1,2) external injection attempt','source_quote':'Original textual clause','source_page':1})
    assert x.status_code==201
    output=client.get(f'/api/tenders/{t}/export.xlsx')
    assert output.status_code==200
    book=load_workbook(io.BytesIO(output.content))
    assert book.active['E8'].data_type!='f'
    assert book.active['E8'].value.startswith("'=SUM")

def ready_bid(client,expires_on=None):
    org=add_org(client);t=add_tender(client,org);r=add_req(client,t)
    ev=evidence(client,org,expires_on=expires_on)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev,'notes':'Checked original evidence'}).status_code==200
    attest(client,t)
    assert verdict(client,t).status_code==201
    assert client.get(f'/api/tenders/{t}').json()['current_decision']=='BID'
    return org,t,r,ev

def stored_decisions(tender):
    with sqlite3.connect(main.DB) as db:
        return db.execute('SELECT * FROM decisions WHERE tender_id=? ORDER BY id',(tender,)).fetchall()

def upload_non_candidate_pdf(client,tender):
    from reportlab.pdfgen.canvas import Canvas
    buf=io.BytesIO();c=Canvas(buf)
    c.drawString(50,750,'This document records administrative contact information.')
    c.showPage();c.save()
    result=client.post(f'/api/tenders/{tender}/upload-pdf',files={'file':('addendum.pdf',buf.getvalue(),'application/pdf')})
    assert result.status_code==201,result.text
    assert result.json()['candidates']==0
    return result.json()['source_id']

def test_zero_candidate_addendum_invalidates_bid_without_rewriting_history(client):
    org,t,r,ev=ready_bid(client)
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    upload_non_candidate_pdf(client,t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['requirements']==before['requirements']
    assert after['fingerprint']==before['fingerprint']
    assert after['source_scope_fingerprint']!=before['source_scope_fingerprint']
    assert after['source_scope_verified'] is False
    assert after['compliance']=='UNRESOLVED'
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert verdict(client,t).status_code==409
    assert stored_decisions(t)==history
    attest(client,t)
    reviewed=client.get(f'/api/tenders/{t}').json()
    assert reviewed['compliance']=='READY_FOR_HUMAN_DECISION'
    assert reviewed['current_decision'] is None
    assert stored_decisions(t)==history
    assert verdict(client,t).status_code==201
    latest=client.get(f'/api/tenders/{t}').json()
    assert latest['current_decision']=='BID'
    assert [d['is_current'] for d in latest['decisions']]==[True,False]
    assert stored_decisions(t)[:len(history)]==history

@pytest.mark.parametrize('field,value',[
    ('name','revised-addendum.pdf'),('sha256','a'*64),('bytes',12345),
    ('pages',2),('uploaded_at','2030-01-01T00:00:00+00:00')
])
def test_source_inventory_mutation_invalidates_bid(client,field,value):
    org,t,r,ev=ready_bid(client)
    source=upload_non_candidate_pdf(client,t)
    attest(client,t);assert verdict(client,t).status_code==201
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    with sqlite3.connect(main.DB) as db:
        db.execute(f'UPDATE sources SET {field}=? WHERE id=?',(value,source))
    after=client.get(f'/api/tenders/{t}').json()
    assert after['source_scope_fingerprint']!=before['source_scope_fingerprint']
    assert after['source_scope_verified'] is False
    assert after['current_decision'] is None
    assert after['compliance']=='UNRESOLVED'
    assert stored_decisions(t)==history

def test_changed_requirement_source_association_invalidates_bid(client):
    org,t,r,ev=ready_bid(client)
    first=upload_non_candidate_pdf(client,t);second=upload_non_candidate_pdf(client,t)
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE requirements SET source_id=? WHERE id=?',(first,r))
    attest(client,t);assert verdict(client,t).status_code==201
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE requirements SET source_id=? WHERE id=?',(second,r))
    after=client.get(f'/api/tenders/{t}').json()
    assert after['source_scope_fingerprint']!=before['source_scope_fingerprint']
    assert after['fingerprint']!=before['fingerprint']
    assert after['source_scope_verified'] is False and after['current_decision'] is None
    assert verdict(client,t).status_code==409
    assert stored_decisions(t)==history

@pytest.mark.parametrize('field,value',[
    ('label','Replacement certificate'),('reference','CERT-REVISED'),
    ('verification_note','Rechecked against a replacement signed original'),
    ('expires_on',(date.today()+timedelta(days=90)).isoformat()),('verified',0)
])
def test_linked_evidence_mutations_invalidate_bid(client,field,value):
    org,t,r,ev=ready_bid(client,expires_on=(date.today()+timedelta(days=60)).isoformat())
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    # Evidence has no update API; change the linked fixture record directly.
    with sqlite3.connect(main.DB) as db:
        db.execute(f'UPDATE evidence SET {field}=? WHERE id=?',(value,ev))
    after=client.get(f'/api/tenders/{t}').json()
    assert after['fingerprint']!=before['fingerprint']
    assert after['source_scope_verified'] is True
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert after['compliance']==('UNRESOLVED' if field=='verified' else 'READY_FOR_HUMAN_DECISION')
    assert stored_decisions(t)==history

def test_linked_evidence_expiry_invalidates_bid_without_a_write(client,monkeypatch):
    expiry=date.today()+timedelta(days=1)
    org,t,r,ev=ready_bid(client,expires_on=expiry.isoformat());history=stored_decisions(t)
    class AfterExpiry(date):
        @classmethod
        def today(cls):
            return expiry+timedelta(days=1)
    monkeypatch.setattr(main,'date',AfterExpiry)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['requirements'][0]['effective_status']=='UNKNOWN'
    assert after['compliance']=='UNRESOLVED' and after['current_decision'] is None
    assert after['decisions'][0]['is_current'] is False
    assert verdict(client,t).status_code==409
    assert stored_decisions(t)==history

def test_unrelated_evidence_and_metrics_preserve_current_bid_and_history(client):
    org,t,r,ev=ready_bid(client)
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    unrelated=evidence(client,org)
    other_org=add_org(client,'Other consulting organization');other=evidence(client,other_org)
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE evidence SET reference=?,verified=0 WHERE id IN (?,?)',('UNRELATED-REVISED',unrelated,other))
    assert client.post(f'/api/tenders/{t}/metrics',json={'paid_bdt':15000,'actual_minutes':75,'notes':'Updated commercial pilot observation'}).status_code==200
    for _ in range(2):
        after=client.get(f'/api/tenders/{t}').json()
        assert after['fingerprint']==before['fingerprint']
        assert after['source_scope_fingerprint']==before['source_scope_fingerprint']
        assert after['current_decision']=='BID' and after['decisions'][0]['is_current'] is True
        assert stored_decisions(t)==history

def test_changed_evidence_mapping_notes_invalidate_bid(client):
    org,t,r,ev=ready_bid(client);history=stored_decisions(t)
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':ev,'notes':'Changed supporting evidence interpretation'}).status_code==200
    after=client.get(f'/api/tenders/{t}').json()
    assert after['compliance']=='READY_FOR_HUMAN_DECISION'
    assert after['current_decision'] is None
    assert stored_decisions(t)==history

@pytest.mark.parametrize('legacy_case',[
    'missing_currency_version','missing_scope_hash','missing_fingerprint','missing_scope_attestation',
    'unsupported_currency_version','null','array','malformed'
])
def test_incomplete_or_legacy_decision_metadata_is_historical(client,legacy_case):
    org,t,r,ev=ready_bid(client)
    with sqlite3.connect(main.DB) as db:
        row=db.execute('SELECT id,decision_snapshot FROM decisions WHERE tender_id=?',(t,)).fetchone()
        past=json.loads(row[1])
        if legacy_case.startswith('missing_'):
            past.pop(legacy_case.removeprefix('missing_'),None)
            text=json.dumps(past)
        elif legacy_case=='unsupported_currency_version':
            past['currency_version']=0;text=json.dumps(past)
        else:
            text={'null':'null','array':'[]','malformed':'{not valid json'}[legacy_case]
        db.execute('UPDATE decisions SET decision_snapshot=? WHERE id=?',(text,row[0]))
    history=stored_decisions(t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['compliance']=='READY_FOR_HUMAN_DECISION'
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert after['decisions'][0]['decision_snapshot']==text
    assert stored_decisions(t)==history

def test_current_bid_requires_current_readiness(client):
    org,t,r,ev=ready_bid(client);history=stored_decisions(t)
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE tenders SET source_scope_fingerprint=NULL WHERE id=?',(t,))
    after=client.get(f'/api/tenders/{t}').json()
    assert after['compliance']=='UNRESOLVED' and after['source_scope_verified'] is False
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert verdict(client,t).status_code==409
    assert stored_decisions(t)==history

@pytest.mark.parametrize('revoke',[False,True])
def test_same_scope_reattestation_does_not_restore_old_bid(client,monkeypatch,revoke):
    monkeypatch.setattr(main,'now',lambda:'2030-01-01T00:00:00+00:00')
    org,t,r,ev=ready_bid(client)
    before=client.get(f'/api/tenders/{t}').json();history=stored_decisions(t)
    if revoke:
        with sqlite3.connect(main.DB) as db:
            db.execute('UPDATE tenders SET source_scope_fingerprint=NULL WHERE id=?',(t,))
        assert client.get(f'/api/tenders/{t}').json()['current_decision'] is None
    attest(client,t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['source_scope_fingerprint']==before['source_scope_fingerprint']
    assert after['source_scope_verified'] is True and after['compliance']=='READY_FOR_HUMAN_DECISION'
    assert after['source_scope_attestation']['at']==before['source_scope_attestation']['at']
    assert after['source_scope_attestation']['event_id']!=before['source_scope_attestation']['event_id']
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert stored_decisions(t)==history
    assert verdict(client,t).status_code==201
    assert client.get(f'/api/tenders/{t}').json()['current_decision']=='BID'
    assert stored_decisions(t)[:len(history)]==history

@pytest.mark.parametrize('field,value',[
    ('source_scope_reviewer','Replacement Human Reviewer'),
    ('source_scope_note','A revised note describing the inspected source set'),
    ('source_scope_at','2030-01-01T00:00:00+00:00')
])
def test_changed_attestation_values_invalidate_bid(client,field,value):
    org,t,r,ev=ready_bid(client);history=stored_decisions(t)
    with sqlite3.connect(main.DB) as db:
        db.execute(f'UPDATE tenders SET {field}=? WHERE id=?',(value,t))
    after=client.get(f'/api/tenders/{t}').json()
    assert after['source_scope_verified'] is True and after['compliance']=='READY_FOR_HUMAN_DECISION'
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert stored_decisions(t)==history

@pytest.mark.parametrize('decision',['HOLD','NO_BID'])
def test_non_bid_decisions_can_remain_current_without_readiness(client,decision):
    org=add_org(client);t=add_tender(client,org);add_req(client,t)
    assert verdict(client,t,decision).status_code==201
    history=stored_decisions(t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['compliance']=='UNRESOLVED'
    assert after['current_decision']==decision and after['decisions'][0]['is_current'] is True
    assert stored_decisions(t)==history

@pytest.mark.parametrize('literal',[
    '=SUM(1,2)','+SUM(1,2)','-SUM(1,2)','@SUM(1,2)',
    ' \t=SUM(1,2)',' \t+SUM(1,2)',' \t-SUM(1,2)',' \t@SUM(1,2)',
    "'=SUM(1,2)",'Ordinary reviewed text'
])
def test_all_untrusted_spreadsheet_fields_are_literals(client,literal):
    import csv
    from openpyxl import load_workbook
    org=add_org(client)
    result=client.post('/api/tenders',json={'organization_id':org,'title':literal})
    assert result.status_code==201,result.text
    t=result.json()['id']
    result=client.post(f'/api/tenders/{t}/requirements',json={'text':literal,'source_page':2,'source_quote':literal})
    assert result.status_code==201,result.text
    r=result.json()['id']
    assert client.put(f'/api/requirements/{r}/review',json={'text':literal,'mandatory':True,'source_page':2,'source_quote':literal}).status_code==200
    result=client.post('/api/evidence',json={'organization_id':org,'label':literal,'reference':'CERT-EXPORT','verified':True,'verification_note':'Checked signed original'})
    assert result.status_code==201,result.text
    assert client.put(f'/api/requirements/{r}/status',json={'status':'VERIFIED','evidence_id':result.json()['id'],'notes':literal}).status_code==200
    result=client.post(f'/api/tenders/{t}/decisions',json={'decision':'HOLD','reviewer':literal,'rationale':literal+' reviewed by a human'})
    assert result.status_code==201,result.text
    book=load_workbook(io.BytesIO(client.get(f'/api/tenders/{t}/export.xlsx').content),data_only=False)
    def escaped(text):
        return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text
    for cell in ['B2','E8','G8','H8','I8']:
        item=book['Compliance Matrix'][cell]
        assert item.data_type!='f'
        assert item.value==escaped(literal.strip())
    for cell,text in [('B2',literal.strip()),('C2',(literal+' reviewed by a human').strip())]:
        item=book['Decision History'][cell]
        assert item.data_type!='f' and item.value==escaped(text)
    assert book['Compliance Matrix']['A8'].value==r
    assert book['Compliance Matrix']['A8'].data_type=='n'
    assert book['Compliance Matrix']['F8'].value==2
    rows=list(csv.DictReader(io.StringIO(client.get(f'/api/tenders/{t}/export.csv').text)))
    for field in ['Requirement','Exact source quote','Evidence','Notes']:
        assert rows[0][field]==escaped(literal.strip())

def test_spreadsheet_export_escapes_stored_leading_whitespace(client):
    from openpyxl import load_workbook
    org,t,r,ev=ready_bid(client)
    literal=' \t=SUM(1,2)'
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE tenders SET title=? WHERE id=?',(literal,t))
        db.execute('UPDATE requirements SET text=?,source_quote=?,notes=? WHERE id=?',(literal,literal,literal,r))
        db.execute('UPDATE evidence SET label=? WHERE id=?',(literal,ev))
        db.execute('UPDATE decisions SET reviewer=?,rationale=? WHERE tender_id=?',(literal,literal,t))
    book=load_workbook(io.BytesIO(client.get(f'/api/tenders/{t}/export.xlsx').content),data_only=False)
    for cell in ['B2','E8','G8','H8','I8']:
        item=book['Compliance Matrix'][cell]
        assert item.value=="'"+literal and item.data_type!='f'
    for cell in ['B2','C2']:
        item=book['Decision History'][cell]
        assert item.value=="'"+literal and item.data_type!='f'

@pytest.mark.parametrize('metadata_case',[
    'float_event_id','string_event_id','null_event_id','missing_event_id',
    'extra_key','missing_reviewer','missing_note','missing_at','missing_scope_hash',
    'boolean_reviewer','null_note','numeric_at','object_scope_hash','array_attestation'
])
def test_attestation_metadata_requires_exact_shape_and_types(client,metadata_case):
    org,t,r,ev=ready_bid(client)
    with sqlite3.connect(main.DB) as db:
        row=db.execute('SELECT id,decision_snapshot FROM decisions WHERE tender_id=? ORDER BY id DESC',(t,)).fetchone()
        past=json.loads(row[1]);metadata=past['scope_attestation']
        if metadata_case=='float_event_id':
            metadata['event_id']=float(metadata['event_id'])
        elif metadata_case=='string_event_id':
            metadata['event_id']=str(metadata['event_id'])
        elif metadata_case=='null_event_id':
            metadata['event_id']=None
        elif metadata_case.startswith('missing_'):
            metadata.pop(metadata_case.removeprefix('missing_'))
        elif metadata_case=='extra_key':
            metadata['unexpected']=None
        elif metadata_case=='boolean_reviewer':
            metadata['reviewer']=True
        elif metadata_case=='null_note':
            metadata['note']=None
        elif metadata_case=='numeric_at':
            metadata['at']=123
        elif metadata_case=='object_scope_hash':
            metadata['scope_hash']={'hash':metadata['scope_hash']}
        else:
            past['scope_attestation']=list(metadata.values())
        db.execute('UPDATE decisions SET decision_snapshot=? WHERE id=?',(json.dumps(past),row[0]))
    history=stored_decisions(t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['compliance']=='READY_FOR_HUMAN_DECISION'
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert stored_decisions(t)==history

def test_boolean_attestation_event_id_cannot_match_integer_one(client):
    org,t,r,ev=ready_bid(client)
    with sqlite3.connect(main.DB) as db:
        # Make the fixture attestation's real event ID 1, where True == 1 in Python.
        db.execute('DELETE FROM audit_events WHERE id=1')
        db.execute("UPDATE audit_events SET id=1 WHERE tender_id=? AND action='human_source_scope_attested'",(t,))
    assert verdict(client,t).status_code==201
    with sqlite3.connect(main.DB) as db:
        row=db.execute('SELECT id,decision_snapshot FROM decisions WHERE tender_id=? ORDER BY id DESC',(t,)).fetchone()
        past=json.loads(row[1]);assert past['scope_attestation']['event_id']==1
        past['scope_attestation']['event_id']=True
        db.execute('UPDATE decisions SET decision_snapshot=? WHERE id=?',(json.dumps(past),row[0]))
    history=stored_decisions(t)
    after=client.get(f'/api/tenders/{t}').json()
    assert after['current_decision'] is None and after['decisions'][0]['is_current'] is False
    assert stored_decisions(t)==history

@pytest.mark.parametrize('literal',[*ERROR_CODES,'00123','123.45','1E10'])
def test_spreadsheet_error_codes_and_numeric_looking_text_are_strings(client,literal):
    from openpyxl import load_workbook
    org,t,r,ev=ready_bid(client)
    with sqlite3.connect(main.DB) as db:
        db.execute('UPDATE tenders SET title=? WHERE id=?',(literal,t))
        db.execute('UPDATE requirements SET text=?,source_quote=?,notes=? WHERE id=?',(literal,literal,literal,r))
        db.execute('UPDATE evidence SET label=? WHERE id=?',(literal,ev))
        db.execute('UPDATE decisions SET reviewer=?,rationale=? WHERE tender_id=?',(literal,literal,t))
    book=load_workbook(io.BytesIO(client.get(f'/api/tenders/{t}/export.xlsx').content),data_only=False)
    for cell in ['B2','E8','G8','H8','I8']:
        item=book['Compliance Matrix'][cell]
        assert item.value==literal and item.data_type=='s'
    for cell in ['B2','C2']:
        item=book['Decision History'][cell]
        assert item.value==literal and item.data_type=='s'
    for sheet in book:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value,str):
                    assert cell.data_type=='s'
    assert book['Compliance Matrix']['A8'].value==r
    assert book['Compliance Matrix']['A8'].data_type=='n'
    assert book['Compliance Matrix']['F8'].value==2
    assert book['Compliance Matrix']['F8'].data_type=='n'
