"""Commercial behavior checks against real SQLite transactions and HTTP routes."""
import csv
import hashlib
import io
import json
import sqlite3
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import main
from app.commercial import CHECKLIST_KEYS, register_routes, report_docx
from app.commercial_schema import (COMMERCIAL_MIGRATION_ID, _LOCAL_COMMERCIAL_STATEMENTS,
                                   initialize_local_commercial)
from app.database import ConfigurationError
from app.workflow_schema import initialize_local_workflows
from test_workflows import (add_org, add_tender, ready_bid, stored_decisions,
                            upload_non_candidate_pdf)


@pytest.fixture
def client(tmp_path, monkeypatch):
    db = tmp_path / 'commercial.db'
    storage = tmp_path / 'uploads'
    storage.mkdir()
    monkeypatch.setattr(main, 'DB', db)
    monkeypatch.setattr(main, 'STORAGE', storage)
    with sqlite3.connect(db) as connection:
        connection.executescript(main.SCHEMA)
        main.initialize_local_extensions(connection)
        initialize_local_workflows(connection)
        initialize_local_commercial(connection)
    if not any(route.path == '/api/commercial/dashboard' for route in main.app.routes):
        register_routes(main.app, main)
    return TestClient(main.app)


def tender(client, organization=None):
    return add_tender(client, organization or add_org(client))


def intake(client, tender_id, **changes):
    body = {'scope': 'Review recorded sources and explain mandatory evidence gaps.',
            'exclusions': 'No bidding or legal eligibility guarantee.',
            'reviewer': 'Owner reviewer', 'quoted_fee_minor': None, 'currency': 'BDT',
            'agreed_inventory_ids': [], 'due_at': None,
            'checklist': {key: True for key in CHECKLIST_KEYS}}
    body.update(changes)
    result = client.put(f'/api/tenders/{tender_id}/service-intake', json=body)
    assert result.status_code == 200, result.text
    return result.json()['intake']


def entry(client, tender_id, key, kind='RECEIPT', amount=10_000, original=None, **changes):
    body = {'kind': kind, 'amount_minor': amount, 'currency': 'BDT',
            'reverses_entry_id': original, 'idempotency_key': key, 'reference': f'Ref {key}', 'note': ''}
    body.update(changes)
    return client.post(f'/api/tenders/{tender_id}/ledger', json=body)


def release(client, tender_id, key='release-0001', **changes):
    body = {'reviewer': 'Owner reviewer', 'delivery_note': 'Reviewed diagnostic recorded for manual delivery.',
            'idempotency_key': key, 'delivery_confirmed_manually': False, 'delivery_reference': ''}
    body.update(changes)
    return client.post(f'/api/tenders/{tender_id}/diagnostic-releases', json=body)


def test_blank_intake_and_quote_do_not_count_as_cash(client):
    t = tender(client)
    assert client.get(f'/api/tenders/{t}/service-intake').json()['intake'] is None
    current = intake(client, t, quoted_fee_minor=125_000)
    assert current['quoted_fee_minor'] == 125_000
    # Historical founder-entered pilot aggregate is kept separate from cash.
    assert client.post(f'/api/tenders/{t}/metrics', json={'paid_bdt': 99_999}).status_code == 200
    dashboard = client.get('/api/commercial/dashboard').json()
    assert dashboard['totals']['net_cash_minor'] == 0
    assert dashboard['paid_assessment_count'] == 0
    assert dashboard['quoted_assessment_count'] == 1
    assert dashboard['unknown_costs'] is True
    assert 'not profit' in dashboard['disclaimer']


@pytest.mark.parametrize('value', [True, False, 1.0, 1.5, '100', 0, -1, None, 9_000_000_000_001])
def test_ledger_requires_exact_positive_integer_minor_units(client, value):
    t = tender(client)
    assert entry(client, t, 'money-0001', amount=value).status_code == 422
    assert client.get(f'/api/tenders/{t}/ledger').json()['entries'] == []


@pytest.mark.parametrize('value', [True, 2.0, '200', -1])
def test_quote_requires_exact_nonnegative_integer_minor_units(client, value):
    t = tender(client)
    response = client.put(f'/api/tenders/{t}/service-intake', json={
        'scope': 'Review source inventory and report evidence gaps.', 'reviewer': 'Owner', 'quoted_fee_minor': value})
    assert response.status_code == 422


def test_nullable_and_zero_quotes_dates_scoped_sources_and_strict_fields(client):
    t = tender(client)
    assert intake(client, t, quoted_fee_minor=0)['quoted_fee_minor'] == 0
    assert intake(client, t, quoted_fee_minor=None)['quoted_fee_minor'] is None
    foreign = upload_non_candidate_pdf(client, tender(client))
    with main.conn() as db:
        inv = db.execute('SELECT id FROM tender_inventory_items WHERE source_id=?', (foreign,)).fetchone()['id']
    body = {'scope': 'Read all supplied sources.', 'reviewer': 'Owner', 'agreed_inventory_ids': [inv]}
    assert client.put(f'/api/tenders/{t}/service-intake', json=body).status_code == 422
    body['agreed_inventory_ids'] = []
    body['due_at'] = '2026-10-20T11:00:00'
    assert client.put(f'/api/tenders/{t}/service-intake', json=body).status_code == 422
    body['due_at'] = '2026-10-20T11:00:00+06:00'
    assert client.put(f'/api/tenders/{t}/service-intake', json=body).json()['intake']['due_at'] == '2026-10-20T05:00:00+00:00'
    body['snapshot'] = {'current_decision': 'BID'}
    assert client.put(f'/api/tenders/{t}/service-intake', json=body).status_code == 422


def test_idempotency_duplicate_reference_and_currency_validation(client):
    t = tender(client)
    first = entry(client, t, 'receipt-0001')
    assert first.status_code == 201
    replay = entry(client, t, 'receipt-0001')
    assert replay.json()['replayed'] is True
    assert replay.json()['entry']['id'] == first.json()['entry']['id']
    assert entry(client, t, 'receipt-0001', amount=10_001).status_code == 409
    assert entry(client, t, 'receipt-0002', reference='Ref receipt-0001').status_code == 409
    assert entry(client, t, 'receipt-0003', currency='USD').status_code == 422
    assert entry(client, t, 'receipt-0004', unexpected=True).status_code == 422
    assert len(client.get(f'/api/tenders/{t}/ledger').json()['entries']) == 1


def test_partial_refunds_and_reversals_reconcile_exactly(client):
    t = tender(client)
    receipt = entry(client, t, 'receipt-0001', amount=10_000).json()['entry']['id']
    cost = entry(client, t, 'cost-000001', kind='DIRECT_COST', amount=2_500).json()['entry']['id']
    refund = entry(client, t, 'refund-0001', kind='REFUND', amount=1_250, original=receipt)
    assert refund.status_code == 201
    assert entry(client, t, 'reverse-0001', kind='REVERSAL', amount=750, original=receipt).status_code == 201
    assert entry(client, t, 'reverse-cost1', kind='REVERSAL', amount=500, original=cost).status_code == 201
    result = client.get(f'/api/tenders/{t}/ledger').json()
    assert result['totals'] == {
        'gross_receipts_minor': 10_000, 'refunds_minor': 1_250, 'receipt_reversals_minor': 750,
        'net_cash_minor': 8_000, 'gross_direct_costs_minor': 2_500, 'cost_reversals_minor': 500,
        'recorded_direct_costs_minor': 2_000, 'contribution_after_recorded_costs_minor': 6_000, 'unknown_costs': True}
    assert next(x for x in result['entries'] if x['id'] == receipt)['remaining_correctable_minor'] == 8_000
    assert entry(client, t, 'refund-over1', kind='REFUND', amount=8_001, original=receipt).status_code == 422
    assert entry(client, t, 'refund-cost1', kind='REFUND', amount=100, original=cost).status_code == 422
    assert entry(client, t, 'reverse-refund1', kind='REVERSAL', amount=100, original=refund.json()['entry']['id']).status_code == 422
    foreign = tender(client)
    assert entry(client, foreign, 'refund-foreign', kind='REFUND', amount=100, original=receipt).status_code == 422
    assert entry(client, t, 'receipt-parent', original=receipt).status_code == 422
    assert entry(client, t, 'refund-no-parent', kind='REFUND', amount=100).status_code == 422


def test_repeat_paid_client_requires_distinct_positive_net_assessments(client):
    org = add_org(client)
    t1, t2 = tender(client, org), tender(client, org)
    r1 = entry(client, t1, 'receipt-first').json()['entry']['id']
    entry(client, t1, 'receipt-extra')
    assert client.get('/api/commercial/dashboard').json()['repeat_paid_organization_count'] == 0
    r2 = entry(client, t2, 'receipt-second').json()['entry']['id']
    assert client.get('/api/commercial/dashboard').json()['repeat_paid_organization_count'] == 1
    entry(client, t2, 'refund-second', kind='REFUND', original=r2)
    dashboard = client.get('/api/commercial/dashboard').json()
    assert dashboard['repeat_paid_organization_count'] == 0
    assert dashboard['paid_assessment_count'] == 1
    assert dashboard['totals']['net_cash_minor'] == 20_000


def test_release_requires_manual_checklist_but_accepts_unresolved_diagnostic(client):
    t = tender(client)
    assert release(client, t).status_code == 422
    intake(client, t, checklist={key: False for key in CHECKLIST_KEYS})
    assert release(client, t).status_code == 422
    intake(client, t)
    response = release(client, t)
    assert response.status_code == 201, response.text
    result = response.json()['release']
    assert result['compliance'] == 'UNRESOLVED'
    assert result['current_decision'] is None
    assert result['warnings']
    assert result['is_current'] is True
    assert client.get(f'/api/tenders/{t}').json()['source_scope_verified'] is False


@pytest.mark.parametrize('decision', ['HOLD', 'NO_BID'])
def test_non_bid_outcomes_are_valid_deliveries_without_eligibility_mutation(client, decision):
    t = tender(client)
    result = client.post(f'/api/tenders/{t}/decisions', json={'decision': decision, 'reviewer': 'Owner', 'rationale': 'Cannot establish requirements until full source inventory exists.'})
    assert result.status_code == 201
    before = stored_decisions(t)
    intake(client, t)
    entry(client, t, 'receipt-outcome')
    released = release(client, t, delivery_confirmed_manually=True, delivery_reference='Client handover register 001')
    assert released.status_code == 201
    assert released.json()['release']['current_decision'] == decision
    assert stored_decisions(t) == before
    assert client.get(f'/api/tenders/{t}').json()['compliance'] == 'UNRESOLVED'


def test_commercial_writes_preserve_current_bid_and_frozen_exports_survive_changes(client):
    org, t, req, ev = ready_bid(client)
    decisions = stored_decisions(t)
    inventory = client.get(f'/api/tenders/{t}').json()['inventory']
    intake(client, t, agreed_inventory_ids=[item['id'] for item in inventory])
    payment = entry(client, t, 'receipt-ready').json()['entry']['id']
    released = release(client, t).json()['release']
    rid = released['id']
    assert client.get(f'/api/tenders/{t}').json()['current_decision'] == 'BID'
    assert stored_decisions(t) == decisions
    base = f'/api/tenders/{t}/diagnostic-releases/{rid}'
    before = {fmt: client.get(f'{base}/export.{fmt}').content for fmt in ('json', 'csv', 'docx')}
    frozen = json.loads(before['json'])
    assert frozen['context']['assessment']['current_decision'] == 'BID'
    assert frozen['context']['assessment']['sources'][0]['sha256']
    assert frozen['context']['evidence_library'][0]['document_hash']
    assert b'private_path' not in before['json']
    assert hashlib.sha256(before['json']).hexdigest() == released['snapshot_sha256']
    assert entry(client, t, 'refund-ready', kind='REFUND', amount=100, original=payment).status_code == 201
    assert client.get(base).json()['is_current'] is False
    upload_non_candidate_pdf(client, t)
    assert client.get(f'/api/tenders/{t}').json()['current_decision'] is None
    for fmt in ('json', 'csv', 'docx'):
        after = client.get(f'{base}/export.{fmt}')
        assert after.content == before[fmt]
        assert after.headers['x-content-sha256'] == hashlib.sha256(after.content).hexdigest()
    assert stored_decisions(t) == decisions
    doc = Document(io.BytesIO(before['docx']))
    assert any('Requirement #' in paragraph.text for paragraph in doc.paragraphs)


def test_release_idempotency_supersession_scope_and_arbitrary_snapshot_rejection(client):
    t = tender(client)
    intake(client, t)
    first = release(client, t).json()['release']
    assert release(client, t).json()['replayed'] is True
    assert release(client, t, delivery_note='Changed payload with the original idempotency key.').status_code == 409
    assert release(client, t, 'release-0002', snapshot={'current_decision': 'BID'}).status_code == 422
    second = release(client, t, 'release-0002').json()['release']
    assert second['release_version'] == 2
    assert client.get(f"/api/tenders/{t}/diagnostic-releases/{first['id']}").json()['is_superseded'] is True
    foreign = tender(client)
    assert client.get(f"/api/tenders/{foreign}/diagnostic-releases/{first['id']}").status_code == 404
    assert client.get(f"/api/tenders/{foreign}/diagnostic-releases/{first['id']}/export.json").status_code == 404
    assert release(client, t, 'release-manual', delivery_confirmed_manually=True).status_code == 422


def test_removed_agreed_source_requires_scope_review_and_preserves_release(client):
    t = tender(client)
    source = upload_non_candidate_pdf(client, t)
    item = client.get(f'/api/tenders/{t}').json()['inventory'][0]
    intake(client, t, agreed_inventory_ids=[item['id']])
    rid = release(client, t).json()['release']['id']
    before = client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}/export.json').content
    assert client.delete(f"/api/inventory/{item['id']}").status_code == 200
    assert release(client, t, 'release-removed').status_code == 422
    assert client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}').json()['is_current'] is False
    assert client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}/export.json').content == before


def test_unlinked_organization_evidence_change_stales_captured_report(client):
    t = tender(client)
    intake(client, t)
    rid = release(client, t).json()['release']['id']
    org = client.get(f'/api/tenders/{t}').json()['tender']['organization_id']
    assert client.post('/api/evidence', json={'organization_id': org, 'label': 'New unrelated evidence', 'reference': 'manual evidence ref'}).status_code == 201
    assert client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}').json()['is_current'] is False


def test_spreadsheet_literals_and_safe_download_names(client):
    org = add_org(client, '=HYPERLINK("evil","client")')
    t = tender(client, org)
    intake(client, t, scope='=HYPERLINK("evil","scope")', exclusions='\t@danger')
    entry(client, t, 'receipt-formula', reference='=DANGER(1)')
    rid = release(client, t, delivery_note='=HYPERLINK("evil","delivery")').json()['release']['id']
    response = client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}/export.csv')
    rows = list(csv.reader(io.StringIO(response.text)))
    assert all(not cell.lstrip().startswith(('=', '+', '-', '@')) for row in rows for cell in row)
    assert any(cell.startswith("'=HYPERLINK") for row in rows for cell in row)
    assert response.headers['content-disposition'] == f'attachment; filename="tenderos_{t}_diagnostic_{rid}.csv"'


def test_ledger_and_releases_are_append_only_even_through_sql(client):
    t = tender(client)
    intake(client, t)
    ledger_id = entry(client, t, 'receipt-locked').json()['entry']['id']
    release_id = release(client, t).json()['release']['id']
    for table, row_id in (('commercial_ledger', ledger_id), ('diagnostic_releases', release_id)):
        for sql in (f'UPDATE {table} SET actor=? WHERE id=?', f'DELETE FROM {table} WHERE id=?'):
            with sqlite3.connect(main.DB) as connection:
                with pytest.raises(sqlite3.IntegrityError, match='immutable'):
                    connection.execute(sql, ('changed', row_id) if sql.startswith('UPDATE') else (row_id,))


def test_v4_fresh_repeat_digest_drift_and_unversioned_refusal(tmp_path):
    with sqlite3.connect(tmp_path / 'fresh.db') as db:
        db.executescript(main.SCHEMA)
        main.initialize_local_extensions(db)
        before = list(db.execute('SELECT migration_id,schema_sha256 FROM tenderos_local_schema_migrations'))
        initialize_local_commercial(db)
        initialize_local_commercial(db)
        assert len(list(db.execute('SELECT * FROM tenderos_local_schema_migrations'))) == len(before) + 1
        assert list(db.execute('SELECT migration_id,schema_sha256 FROM tenderos_local_schema_migrations WHERE migration_id<>?', (COMMERCIAL_MIGRATION_ID,))) == before
        db.execute('UPDATE tenderos_local_schema_migrations SET schema_sha256=? WHERE migration_id=?', ('0' * 64, COMMERCIAL_MIGRATION_ID))
        with pytest.raises(ConfigurationError, match='digest differs'):
            initialize_local_commercial(db)
    with sqlite3.connect(tmp_path / 'unversioned.db') as db:
        db.executescript(main.SCHEMA)
        db.execute(_LOCAL_COMMERCIAL_STATEMENTS[0])
        with pytest.raises(ConfigurationError, match='Unversioned'):
            initialize_local_commercial(db)
    with pytest.raises(ConfigurationError, match='SQLite'):
        initialize_local_commercial(SimpleNamespace())


def test_nonhex_frozen_hash_rejected_by_local_schema(client):
    t = tender(client)
    with sqlite3.connect(main.DB) as db:
        with pytest.raises(sqlite3.IntegrityError):
            db.execute('''INSERT INTO diagnostic_releases(tender_id,release_version,idempotency_key,request_hash,
                snapshot_json,snapshot_sha256,currency_fingerprint,template_version,actor,reviewer,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)''', (t, 1, 'bad-hash1', '0' * 64, '{}', 'Z' * 64, '0' * 64, 'v1', 'owner', 'owner', 'now'))


def test_frozen_integrity_check_fails_closed_after_out_of_band_corruption(client):
    t = tender(client)
    intake(client, t)
    rid = release(client, t).json()['release']['id']
    with sqlite3.connect(main.DB) as db:
        db.execute('DROP TRIGGER diagnostic_releases_no_update')
        db.execute('UPDATE diagnostic_releases SET snapshot_json=? WHERE id=?', ('{}', rid))
    assert client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}').status_code == 409
    assert client.get(f'/api/tenders/{t}/diagnostic-releases/{rid}/export.json').status_code == 409


def test_word_xml_controls_have_visible_markers_without_changing_snapshot():
    payload = {
        'diagnostic_identifier': 'Diagnostic\x00one', 'captured_at': '2026-10-10T07:00:00+00:00',
        'template_version': 'v1', 'reviewer': 'Reviewer\x0bName', 'actor': 'Owner',
        'context_sha256': '0' * 64, 'limitations': 'Diagnostic limits\x01', 'warnings': ['Warning\x02'],
        'delivery_note': 'Delivery\x03 note', 'delivery_confirmed_manually': False, 'delivery_reference': 'ref\x04',
        'financial_totals': {'net_cash_minor': 0},
        'context': {
            'organization': {'name': 'Organization\x05'}, 'evidence_library': [],
            'intake': {'scope': 'Review supplied documents.\x00', 'exclusions': 'No bids\x06', 'quoted_fee_minor': None},
            'assessment': {'tender': {'title': 'Tender\x07'}, 'current_decision': None, 'compliance': 'UNRESOLVED',
                'inventory': [], 'sources': [{'id': 1, 'name': 'Source\x08', 'pages': 1, 'sha256': 'a' * 64}],
                'requirements': [], 'blockers': [], 'unresolved': [], 'decisions': []}}}
    before = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    generated = report_docx(payload)
    document = Document(io.BytesIO(generated))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    for marker in ('[U+0000]', '[U+0001]', '[U+0002]', '[U+0003]', '[U+0004]',
                   '[U+0005]', '[U+0006]', '[U+0007]', '[U+0008]', '[U+000B]'):
        assert marker in text
    assert 'unsupported by the Word XML format' in text
    assert document.core_properties.author == 'Reviewer[U+000B]Name'
    assert json.dumps(payload, ensure_ascii=False, sort_keys=True) == before
    assert report_docx(payload) == generated
