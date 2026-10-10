"""Human-recorded expert diagnostics, with frozen reports and an append-only ledger.

This module has no model, payment, email or network client. Route dependencies
are injected by main, so its security middleware and transaction factory apply.
"""
import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from typing import Annotated, Literal, Optional

from docx import Document
from fastapi import HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field, StrictBool, field_validator

TEMPLATE_VERSION = 'tenderos-expert-diagnostic-v1'
MAX_RELEASE_BYTES = 4 * 1024 * 1024
CHECKLIST_KEYS = ('intake_complete', 'source_scope_reviewed', 'evidence_gaps_reviewed',
                  'decision_limitations_reviewed')
MONEY_DISCLAIMER = ('Owner-entered offline payment and cost records; not bank-verified. '
                    'Quotes and legacy pilot amounts are not collected revenue. '
                    'Contribution subtracts recorded direct costs only; overhead and other costs are unknown. It is not profit.')
DIAGNOSTIC_DISCLAIMER = ('Expert diagnostic of the recorded document set, not legal qualification, '
                         'a win probability, a bid submission, or a guarantee of eligibility. '
                         'Unresolved findings and HOLD/NO_BID outcomes may be valid diagnostic deliverables. '
                         'Release creates a frozen report; delivery is only manually confirmed when explicitly recorded.')
MinorAmount = Annotated[int, Field(strict=True, ge=1, le=9_000_000_000_000)]
PositiveId = Annotated[int, Field(strict=True, ge=1)]


class StrictInput(BaseModel):
    model_config = ConfigDict(extra='forbid')


class DeliveryChecklist(StrictInput):
    intake_complete: StrictBool = False
    source_scope_reviewed: StrictBool = False
    evidence_gaps_reviewed: StrictBool = False
    decision_limitations_reviewed: StrictBool = False


class IntakeIn(StrictInput):
    scope: str = Field(min_length=10, max_length=6000)
    exclusions: str = Field(default='', max_length=4000)
    agreed_inventory_ids: list[PositiveId] = Field(default_factory=list, max_length=100)
    due_at: Optional[datetime] = None
    quoted_fee_minor: Optional[Annotated[int, Field(strict=True, ge=0, le=9_000_000_000_000)]] = None
    currency: Literal['BDT'] = 'BDT'
    reviewer: str = Field(min_length=2, max_length=100)
    checklist: DeliveryChecklist = Field(default_factory=DeliveryChecklist)

    @field_validator('scope', 'reviewer')
    @classmethod
    def meaningful_text(cls, value, info):
        value = value.strip()
        if len(value) < (10 if info.field_name == 'scope' else 2):
            raise ValueError('A meaningful scope/reviewer is required')
        return value

    @field_validator('due_at')
    @classmethod
    def timezone_required(cls, value):
        if value is not None and (value.tzinfo is None or value.utcoffset() is None):
            raise ValueError('The due date must include a timezone')
        return value.astimezone(timezone.utc) if value else None

    @field_validator('agreed_inventory_ids')
    @classmethod
    def unique_ids(cls, value):
        if len(value) != len(set(value)):
            raise ValueError('Agreed document IDs must be unique')
        return sorted(value)


class LedgerIn(StrictInput):
    kind: Literal['RECEIPT', 'DIRECT_COST', 'REFUND', 'REVERSAL']
    amount_minor: MinorAmount
    currency: Literal['BDT'] = 'BDT'
    reverses_entry_id: Optional[PositiveId] = None
    idempotency_key: str = Field(min_length=8, max_length=100, pattern=r'^[A-Za-z0-9_.:-]+$')
    reference: str = Field(min_length=2, max_length=200)
    note: str = Field(default='', max_length=2000)

    @field_validator('reference')
    @classmethod
    def reference_required(cls, value):
        value = value.strip()
        if len(value) < 2:
            raise ValueError('A manual payment/cost reference is required')
        return value


class ReleaseIn(StrictInput):
    reviewer: str = Field(min_length=2, max_length=100)
    delivery_note: str = Field(min_length=10, max_length=4000)
    idempotency_key: str = Field(min_length=8, max_length=100, pattern=r'^[A-Za-z0-9_.:-]+$')
    delivery_confirmed_manually: StrictBool = False
    delivery_reference: str = Field(default='', max_length=200)

    @field_validator('reviewer', 'delivery_note')
    @classmethod
    def meaningful_text(cls, value, info):
        value = value.strip()
        if len(value) < (10 if info.field_name == 'delivery_note' else 2):
            raise ValueError('A reviewer and meaningful delivery note are required')
        return value


def canonical_json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def sha256_json(value):
    return hashlib.sha256(canonical_json(value).encode('utf-8')).hexdigest()


def intake_view(row):
    if row is None:
        return None
    out = dict(row)
    out['agreed_inventory_ids'] = json.loads(out.pop('agreed_inventory_ids_json'))
    out['checklist'] = json.loads(out.pop('checklist_json'))
    return out


def read_intake(db, tender_id):
    return intake_view(db.execute('SELECT * FROM service_intakes WHERE tender_id=?', (tender_id,)).fetchone())


def ledger_entries(db, tender_id):
    entries = [dict(row) for row in db.execute('SELECT * FROM commercial_ledger WHERE tender_id=? ORDER BY id', (tender_id,))]
    corrected = {}
    for entry in entries:
        if entry['reverses_entry_id'] is not None:
            original_id = entry['reverses_entry_id']
            corrected[original_id] = corrected.get(original_id, 0) + entry['amount_minor']
    for entry in entries:
        entry.pop('request_hash')
        entry['remaining_correctable_minor'] = (entry['amount_minor'] - corrected.get(entry['id'], 0)
                                                if entry['kind'] in ('RECEIPT', 'DIRECT_COST') else None)
    return entries


def ledger_totals(entries):
    originals = {entry['id']: entry for entry in entries}
    gross_receipts = sum(x['amount_minor'] for x in entries if x['kind'] == 'RECEIPT')
    refunds = sum(x['amount_minor'] for x in entries if x['kind'] == 'REFUND')
    receipt_reversals = sum(x['amount_minor'] for x in entries if x['kind'] == 'REVERSAL'
                           and originals[x['reverses_entry_id']]['kind'] == 'RECEIPT')
    gross_costs = sum(x['amount_minor'] for x in entries if x['kind'] == 'DIRECT_COST')
    cost_reversals = sum(x['amount_minor'] for x in entries if x['kind'] == 'REVERSAL'
                        and originals[x['reverses_entry_id']]['kind'] == 'DIRECT_COST')
    net_cash = gross_receipts - refunds - receipt_reversals
    costs = gross_costs - cost_reversals
    return {'gross_receipts_minor': gross_receipts, 'refunds_minor': refunds,
            'receipt_reversals_minor': receipt_reversals, 'net_cash_minor': net_cash,
            'gross_direct_costs_minor': gross_costs, 'cost_reversals_minor': cost_reversals,
            'recorded_direct_costs_minor': costs, 'contribution_after_recorded_costs_minor': net_cash - costs,
            'unknown_costs': True}


def lock_tender(db, domain, tender_id):
    # Acquire the writer lock before any read/check. PostgreSQL serializable
    # conflicts are handled by the shared connection factory; SQLite avoids
    # deferred read-to-write lock upgrades. No domain values are changed.
    cursor = db.execute('UPDATE tenders SET id=id WHERE id=?', (tender_id,))
    if cursor.rowcount != 1:
        raise HTTPException(404, 'Tender record not found')
    return domain.require(db, 'tenders', tender_id)


def report_context(db, domain, tender_id):
    assessment = domain.get_snapshot(db, tender_id)
    organization_id = assessment['tender']['organization_id']
    organization = dict(domain.require(db, 'organizations', organization_id))
    evidence = [domain.evidence_view(db, row) for row in db.execute(
        'SELECT * FROM evidence WHERE organization_id=? ORDER BY id', (organization_id,))]
    # evidence_view/source lists intentionally expose metadata/hash, never private_path.
    return {'assessment': assessment, 'organization': organization,
            'evidence_library': evidence, 'intake': read_intake(db, tender_id),
            'ledger': ledger_entries(db, tender_id)}


def diagnostic_warnings(context):
    snap = context['assessment']
    warnings = []
    if not snap['inventory_complete']:
        warnings.append('The tender source inventory is incomplete or unconfirmed.')
    if not snap['source_scope_verified']:
        warnings.append('A current human attestation of the full source set is absent.')
    if snap['compliance'] != 'READY_FOR_HUMAN_DECISION':
        warnings.append(f"Assessment remains {snap['compliance']}; unresolved findings are included, not cleared by release.")
    if snap['current_decision'] is None:
        warnings.append('No current authorized human BID/NO_BID/HOLD decision is recorded.')
    elif snap['current_decision'] in ('HOLD', 'NO_BID'):
        warnings.append(f"Current human outcome is {snap['current_decision']}.")
    if context['intake']['quoted_fee_minor'] is None:
        warnings.append('No fee quote is recorded; this does not establish collected revenue.')
    return warnings


def release_row(db, tender_id, release_id):
    row = db.execute('SELECT * FROM diagnostic_releases WHERE tender_id=? AND id=?', (tender_id, release_id)).fetchone()
    if not row:
        raise HTTPException(404, 'Diagnostic release not found for this tender')
    # Detect storage mutation without generating a report from current values.
    if hashlib.sha256(row['snapshot_json'].encode('utf-8')).hexdigest() != row['snapshot_sha256']:
        raise HTTPException(409, 'Frozen diagnostic integrity check failed')
    return row


def release_summary(row, current_fingerprint, newest_version):
    payload = json.loads(row['snapshot_json'])
    snap = payload['context']['assessment']
    return {'id': row['id'], 'tender_id': row['tender_id'], 'release_version': row['release_version'],
            'created_at': row['created_at'], 'reviewer': row['reviewer'], 'actor': row['actor'],
            'template_version': row['template_version'], 'snapshot_sha256': row['snapshot_sha256'],
            'is_current': row['currency_fingerprint'] == current_fingerprint,
            'is_superseded': row['release_version'] < newest_version,
            'delivery_confirmed_manually': payload['delivery_confirmed_manually'],
            'delivery_reference': payload['delivery_reference'], 'compliance': snap['compliance'],
            'current_decision': snap['current_decision'], 'warnings': payload['warnings']}


def newest_release_version(db, tender_id):
    row = db.execute('SELECT MAX(release_version) AS version FROM diagnostic_releases WHERE tender_id=?', (tender_id,)).fetchone()
    return row['version'] or 0


def literal(value):
    # Prefix after whitespace as well; numbers remain integer values.
    return "'" + value if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')) else value


def report_csv(payload):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(['Section', 'Field', 'Value'])

    def add(section, field, value):
        if isinstance(value, (dict, list)):
            value = canonical_json(value)
        writer.writerow([literal(section), literal(field), literal(value)])

    context = payload['context']
    for key in ('diagnostic_identifier', 'captured_at', 'template_version', 'reviewer', 'actor',
                'delivery_note', 'delivery_confirmed_manually', 'delivery_reference', 'warnings', 'limitations'):
        add('Diagnostic', key, payload[key])
    for key, value in context['intake'].items():
        add('Intake', key, value)
    for key in ('compliance', 'current_decision', 'mandatory_total', 'mandatory_verified',
                'source_scope_verified', 'inventory_complete', 'source_scope_fingerprint', 'fingerprint'):
        add('Assessment', key, context['assessment'][key])
    for section, records in (('Sources', context['assessment']['sources']),
                             ('Inventory', context['assessment']['inventory']),
                             ('Requirements', context['assessment']['requirements']),
                             ('Evidence', context['evidence_library']),
                             ('Decisions', context['assessment']['decisions']), ('Ledger', context['ledger'])):
        for record in records:
            for key, value in record.items():
                add(f"{section} #{record['id']}", key, value)
    for key, value in payload['financial_totals'].items():
        add('Financial totals (BDT minor units)', key, value)
    return out.getvalue().encode('utf-8')


def word_text(value):
    """Render XML 1.0 unsupported characters visibly without changing the snapshot."""
    return ''.join(character if (ord(character) in (9, 10, 13)
                   or 0x20 <= ord(character) <= 0xD7FF
                   or 0xE000 <= ord(character) <= 0xFFFD
                   or 0x10000 <= ord(character) <= 0x10FFFF)
                   else f'[U+{ord(character):04X}]' for character in str(value))


class _WordDocument:
    """Keep every heading and paragraph on the same XML-safe rendering path."""
    def __init__(self):
        self._document = Document()

    @property
    def core_properties(self):
        return self._document.core_properties

    def add_heading(self, text, level=1):
        return self._document.add_heading(word_text(text), level=level)

    def add_paragraph(self, text, style=None):
        return self._document.add_paragraph(word_text(text), style=style)

    def save(self, target):
        self._document.save(target)


def report_docx(payload):
    document = _WordDocument()
    document.core_properties.title = 'TenderOS expert tender diagnostic'
    document.core_properties.author = word_text(payload['reviewer'])
    captured = datetime.fromisoformat(payload['captured_at']).replace(tzinfo=None)
    document.core_properties.created = captured
    document.core_properties.modified = captured
    context = payload['context']
    snap = context['assessment']
    document.add_heading('TenderOS expert tender diagnostic', 0)
    document.add_paragraph(payload['diagnostic_identifier'])
    document.add_paragraph(f"{context['organization']['name']} — {snap['tender']['title']}")
    document.add_paragraph(f"Captured {payload['captured_at']} · Reviewer: {payload['reviewer']}")
    document.add_heading('Scope and limitations', level=1)
    document.add_paragraph(context['intake']['scope'])
    document.add_paragraph('Exclusions: ' + (context['intake']['exclusions'] or 'No additional exclusions recorded.'))
    document.add_paragraph(payload['limitations'])
    document.add_paragraph('Characters unsupported by the Word XML format are represented with visible '
                           '[U+XXXX] markers. The frozen JSON and CSV exports preserve the original recorded text.')
    document.add_paragraph(f"Outcome: {snap['current_decision'] or 'No current human decision'}; assessment: {snap['compliance']}")
    for warning in payload['warnings']:
        document.add_paragraph(warning, style='List Bullet')
    document.add_heading('Source inventory and provenance', level=1)
    for item in snap['inventory']:
        document.add_paragraph(f"Inventory #{item['id']}: {item['title']} · {item['availability']} · "
                               f"{item['language']} · {item['version_label']} · source #{item['source_id']}")
    for source in snap['sources']:
        document.add_paragraph(f"Source #{source['id']}: {source['name']} · {source['pages']} pages · SHA-256 {source['sha256']}")
    document.add_heading('Requirement and evidence matrix', level=1)
    for req in snap['requirements']:
        document.add_heading(f"Requirement #{req['id']}", level=2)
        document.add_paragraph(req['text'])
        classification = 'mandatory' if req['mandatory'] == 1 else ('optional' if req['mandatory'] == 0 else 'unclassified')
        document.add_paragraph(f"{classification}; reviewed: {bool(req['reviewed'])}; effective status: {req['effective_status']}")
        document.add_paragraph(f"Source #{req['source_id']} page {req['source_page']}: {req['source_quote'] or ''}")
        document.add_paragraph(f"Source SHA-256: {req['source_sha256'] or 'missing'}")
        document.add_paragraph(f"Evidence: {req['evidence_label'] or 'No linked evidence'}; SHA-256: {req['evidence_document_hash'] or 'missing'}")
        if req.get('notes') or req.get('evidence_issue'):
            document.add_paragraph(req.get('evidence_issue') or req['notes'])
    document.add_heading('Next review actions', level=1)
    for item in snap['blockers']:
        document.add_paragraph(f"Mandatory blocker #{item['requirement_id']}: {item['text']}", style='List Bullet')
    for item in snap['unresolved']:
        document.add_paragraph(canonical_json(item), style='List Bullet')
    document.add_heading('Evidence library at capture', level=1)
    for evidence in context['evidence_library']:
        document.add_paragraph(f"Evidence #{evidence['id']}: {evidence['label']} · expires: {evidence['expires_on'] or 'not recorded'} · "
                               f"current verification: {evidence['verification_current']} · SHA-256: {evidence['document_hash'] or 'missing'}")
    document.add_heading('Human decisions at capture', level=1)
    for decision in snap['decisions']:
        document.add_paragraph(f"{decision['decision']} · {decision['reviewer']} · {decision['created_at']} · current: {decision['is_current']}")
        document.add_paragraph(decision['rationale'])
    document.add_heading('Recorded commercial observations', level=1)
    document.add_paragraph(MONEY_DISCLAIMER)
    document.add_paragraph('All monetary values below are integer BDT minor units (100 = BDT 1).')
    document.add_paragraph(f"Quoted fee: {context['intake']['quoted_fee_minor'] if context['intake']['quoted_fee_minor'] is not None else 'not recorded'}")
    for key, value in payload['financial_totals'].items():
        document.add_paragraph(f'{key}: {value}')
    document.add_paragraph(payload['delivery_note'])
    document.add_paragraph(f"Manually confirmed delivery: {payload['delivery_confirmed_manually']}; reference: {payload['delivery_reference'] or 'none'}")
    document.add_heading('Release trace', level=1)
    document.add_paragraph(f"Template: {payload['template_version']} · Captured context SHA-256: {payload['context_sha256']}")
    document.add_paragraph('The frozen JSON export preserves the complete captured metadata and history.')
    raw = io.BytesIO()
    document.save(raw)
    # Normalize ZIP member timestamps so repeat exports of the frozen report
    # have identical bytes and a meaningful downloadable content checksum.
    normalized = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(normalized, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            member = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            member.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(member, source.read(name))
    return normalized.getvalue()


def register_routes(app, domain):
    """Attach the commercial API to the existing authenticated owner workspace."""
    @app.get('/api/tenders/{tender_id}/service-intake')
    def get_intake(tender_id: int):
        with domain.conn() as db:
            domain.require(db, 'tenders', tender_id)
            intake = read_intake(db, tender_id)
        return {'exists': intake is not None, 'intake': intake, 'checklist_keys': CHECKLIST_KEYS,
                'currency': 'BDT', 'money_unit': 'minor_units', 'disclaimer': DIAGNOSTIC_DISCLAIMER}

    @app.put('/api/tenders/{tender_id}/service-intake')
    def put_intake(tender_id: int, item: IntakeIn):
        with domain.conn() as db:
            lock_tender(db, domain, tender_id)
            for inventory_id in item.agreed_inventory_ids:
                row = db.execute('SELECT id FROM tender_inventory_items WHERE id=? AND tender_id=?', (inventory_id, tender_id)).fetchone()
                if not row:
                    raise HTTPException(422, 'An agreed document does not belong to this tender inventory')
            actor = domain.authenticated_actor(item.reviewer)
            old = read_intake(db, tender_id)
            values = (item.scope, item.exclusions, canonical_json(item.agreed_inventory_ids),
                      item.due_at.isoformat() if item.due_at else None, item.quoted_fee_minor,
                      item.currency, item.reviewer, canonical_json(item.checklist.model_dump()), actor, domain.now())
            if old:
                db.execute('''UPDATE service_intakes SET scope=?,exclusions=?,agreed_inventory_ids_json=?,due_at=?,
                    quoted_fee_minor=?,currency=?,reviewer=?,checklist_json=?,actor=?,updated_at=? WHERE tender_id=?''', values + (tender_id,))
            else:
                db.execute('''INSERT INTO service_intakes(scope,exclusions,agreed_inventory_ids_json,due_at,
                    quoted_fee_minor,currency,reviewer,checklist_json,actor,updated_at,tender_id) VALUES(?,?,?,?,?,?,?,?,?,?,?)''', values + (tender_id,))
            domain.log(db, tender_id, 'human_service_intake_recorded',
                       {'scope_sha256': sha256_json(item.model_dump(mode='json')), 'quoted_fee_minor': item.quoted_fee_minor}, actor)
            intake = read_intake(db, tender_id)
        return {'exists': True, 'intake': intake, 'checklist_keys': CHECKLIST_KEYS,
                'currency': 'BDT', 'money_unit': 'minor_units', 'disclaimer': DIAGNOSTIC_DISCLAIMER}

    @app.get('/api/tenders/{tender_id}/ledger')
    def get_ledger(tender_id: int):
        with domain.conn() as db:
            domain.require(db, 'tenders', tender_id)
            entries = ledger_entries(db, tender_id)
        return {'entries': entries, 'totals': ledger_totals(entries), 'currency': 'BDT', 'disclaimer': MONEY_DISCLAIMER}

    @app.post('/api/tenders/{tender_id}/ledger', status_code=201)
    def post_ledger(tender_id: int, item: LedgerIn):
        request_hash = sha256_json(item.model_dump())
        with domain.conn() as db:
            lock_tender(db, domain, tender_id)
            prior = db.execute('SELECT * FROM commercial_ledger WHERE tender_id=? AND idempotency_key=?',
                               (tender_id, item.idempotency_key)).fetchone()
            if prior:
                if prior['request_hash'] != request_hash:
                    raise HTTPException(409, 'Idempotency key was used with different ledger values')
                return {'entry': next(x for x in ledger_entries(db, tender_id) if x['id'] == prior['id']), 'replayed': True}
            duplicate_reference = db.execute('SELECT id FROM commercial_ledger WHERE tender_id=? AND kind=? AND reference=?',
                                             (tender_id, item.kind, item.reference)).fetchone()
            if duplicate_reference:
                raise HTTPException(409, 'This entry kind and manual reference are already recorded; use the original idempotency key to retry')
            if item.kind in ('RECEIPT', 'DIRECT_COST'):
                if item.reverses_entry_id is not None:
                    raise HTTPException(422, 'Original receipt/cost entries cannot reference a correction')
            else:
                original = db.execute('SELECT * FROM commercial_ledger WHERE tender_id=? AND id=?',
                                      (tender_id, item.reverses_entry_id)).fetchone()
                if original is None or original['kind'] not in ('RECEIPT', 'DIRECT_COST'):
                    raise HTTPException(422, 'Correction must reference an original entry in this tender')
                if item.kind == 'REFUND' and original['kind'] != 'RECEIPT':
                    raise HTTPException(422, 'A refund must reference an original receipt')
                corrected = db.execute('SELECT COALESCE(SUM(amount_minor),0) AS amount FROM commercial_ledger WHERE tender_id=? AND reverses_entry_id=?',
                                       (tender_id, original['id'])).fetchone()['amount']
                if corrected + item.amount_minor > original['amount_minor']:
                    raise HTTPException(422, 'Total refunds/reversals exceed the original amount')
            actor = domain.authenticated_actor()
            cursor = db.execute('''INSERT INTO commercial_ledger(tender_id,kind,amount_minor,currency,reverses_entry_id,
                idempotency_key,request_hash,reference,note,actor,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)''',
                (tender_id, item.kind, item.amount_minor, item.currency, item.reverses_entry_id,
                 item.idempotency_key, request_hash, item.reference, item.note, actor, domain.now()))
            domain.log(db, tender_id, 'human_commercial_ledger_recorded',
                       {'entry_id': cursor.lastrowid, 'kind': item.kind, 'amount_minor': item.amount_minor,
                        'reverses_entry_id': item.reverses_entry_id}, actor)
            entry = next(x for x in ledger_entries(db, tender_id) if x['id'] == cursor.lastrowid)
        return {'entry': entry, 'replayed': False}

    @app.get('/api/tenders/{tender_id}/diagnostic-releases')
    def get_releases(tender_id: int):
        with domain.conn() as db:
            fingerprint = sha256_json(report_context(db, domain, tender_id))
            latest = newest_release_version(db, tender_id)
            releases = [release_summary(row, fingerprint, latest) for row in db.execute(
                'SELECT * FROM diagnostic_releases WHERE tender_id=? ORDER BY release_version DESC', (tender_id,))]
        return {'releases': releases, 'template_version': TEMPLATE_VERSION, 'disclaimer': DIAGNOSTIC_DISCLAIMER}

    @app.post('/api/tenders/{tender_id}/diagnostic-releases', status_code=201)
    def post_release(tender_id: int, item: ReleaseIn):
        if item.delivery_confirmed_manually and len(item.delivery_reference.strip()) < 2:
            raise HTTPException(422, 'Manual delivery confirmation requires a reference')
        request_hash = sha256_json(item.model_dump())
        with domain.conn() as db:
            lock_tender(db, domain, tender_id)
            prior = db.execute('SELECT * FROM diagnostic_releases WHERE tender_id=? AND idempotency_key=?',
                               (tender_id, item.idempotency_key)).fetchone()
            context = report_context(db, domain, tender_id)
            fingerprint = sha256_json(context)
            latest = newest_release_version(db, tender_id)
            if prior:
                if prior['request_hash'] != request_hash:
                    raise HTTPException(409, 'Idempotency key was used with a different release request')
                return {'release': release_summary(prior, fingerprint, latest), 'replayed': True}
            intake = context['intake']
            if not intake or not all(intake['checklist'].get(key) is True for key in CHECKLIST_KEYS):
                raise HTTPException(422, 'Record intake and acknowledge all diagnostic delivery checklist items first')
            inventory_ids = {x['id'] for x in context['assessment']['inventory']}
            if not set(intake['agreed_inventory_ids']).issubset(inventory_ids):
                raise HTTPException(422, 'An agreed source was removed; review and update the intake scope')
            actor = domain.authenticated_actor(item.reviewer)
            created_at = domain.now()
            version = latest + 1
            payload = {'diagnostic_identifier': f'TENDER-{tender_id}-DIAGNOSTIC-{version}',
                       'captured_at': created_at, 'template_version': TEMPLATE_VERSION,
                       'reviewer': item.reviewer, 'actor': actor, 'context': context,
                       'context_sha256': fingerprint, 'financial_totals': ledger_totals(context['ledger']),
                       'warnings': diagnostic_warnings(context), 'limitations': DIAGNOSTIC_DISCLAIMER,
                       'delivery_note': item.delivery_note, 'delivery_confirmed_manually': item.delivery_confirmed_manually,
                       'delivery_reference': item.delivery_reference.strip()}
            text = canonical_json(payload)
            if len(text.encode('utf-8')) > MAX_RELEASE_BYTES:
                raise HTTPException(422, 'Diagnostic exceeds the frozen report size limit; reduce the assessment scope')
            digest = hashlib.sha256(text.encode('utf-8')).hexdigest()
            cursor = db.execute('''INSERT INTO diagnostic_releases(tender_id,release_version,idempotency_key,request_hash,
                snapshot_json,snapshot_sha256,currency_fingerprint,template_version,actor,reviewer,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)''',
                (tender_id, version, item.idempotency_key, request_hash, text, digest, fingerprint,
                 TEMPLATE_VERSION, actor, item.reviewer, created_at))
            domain.log(db, tender_id, 'human_diagnostic_released',
                       {'release_id': cursor.lastrowid, 'release_version': version, 'snapshot_sha256': digest,
                        'delivery_confirmed_manually': item.delivery_confirmed_manually}, actor)
            row = release_row(db, tender_id, cursor.lastrowid)
            result = release_summary(row, fingerprint, version)
        return {'release': result, 'replayed': False}

    @app.get('/api/tenders/{tender_id}/diagnostic-releases/{release_id}')
    def get_release(tender_id: int, release_id: int):
        with domain.conn() as db:
            row = release_row(db, tender_id, release_id)
            fingerprint = sha256_json(report_context(db, domain, tender_id))
            result = release_summary(row, fingerprint, newest_release_version(db, tender_id))
            result['snapshot'] = json.loads(row['snapshot_json'])
        return result

    @app.get('/api/tenders/{tender_id}/diagnostic-releases/{release_id}/export.{format}')
    def export_release(tender_id: int, release_id: int, format: str):
        if format not in ('json', 'csv', 'docx'):
            raise HTTPException(404, 'Diagnostic export format not available')
        with domain.conn() as db:
            row = release_row(db, tender_id, release_id)
            payload = json.loads(row['snapshot_json'])
        if format == 'json':
            content = row['snapshot_json'].encode('utf-8')
            media_type = 'application/json'
        elif format == 'csv':
            content = report_csv(payload)
            media_type = 'text/csv; charset=utf-8'
        else:
            content = report_docx(payload)
            media_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        return Response(content, media_type=media_type, headers={
            'Content-Disposition': f'attachment; filename="tenderos_{tender_id}_diagnostic_{release_id}.{format}"',
            'X-Content-SHA256': hashlib.sha256(content).hexdigest(),
            'X-Diagnostic-Snapshot-SHA256': row['snapshot_sha256']})

    @app.get('/api/commercial/dashboard')
    def commercial_dashboard():
        with domain.conn() as db:
            assessments = []
            all_entries = []
            org_paid = {}
            release_count = manually_delivered = quoted = paid = 0
            for tender in db.execute('''SELECT tenders.id,tenders.organization_id,tenders.title,organizations.name AS organization_name
                FROM tenders JOIN organizations ON organizations.id=tenders.organization_id ORDER BY tenders.id'''):
                entries = ledger_entries(db, tender['id'])
                # Ledger IDs are globally unique, so totals can resolve corrections
                # safely even when aggregating distinct scoped assessments.
                all_entries.extend(entries)
                totals = ledger_totals(entries)
                intake = read_intake(db, tender['id'])
                if intake and intake['quoted_fee_minor'] is not None:
                    quoted += 1
                if totals['net_cash_minor'] > 0:
                    paid += 1
                    org_paid[tender['organization_id']] = org_paid.get(tender['organization_id'], 0) + 1
                releases = list(db.execute('SELECT snapshot_json FROM diagnostic_releases WHERE tender_id=?', (tender['id'],)))
                release_count += len(releases)
                manually_delivered += sum(1 for r in releases if json.loads(r['snapshot_json'])['delivery_confirmed_manually'])
                assessments.append({**dict(tender), 'tender_id': tender['id'],
                                    'quoted_fee_minor': intake['quoted_fee_minor'] if intake else None,
                                    'net_cash_minor': totals['net_cash_minor'],
                                    'recorded_direct_costs_minor': totals['recorded_direct_costs_minor'],
                                    'contribution_after_recorded_costs_minor': totals['contribution_after_recorded_costs_minor'],
                                    'release_count': len(releases)})
        return {'currency': 'BDT', 'money_unit': 'minor_units', 'totals': ledger_totals(all_entries),
                'assessment_count': len(assessments), 'paid_assessment_count': paid,
                'quoted_assessment_count': quoted, 'repeat_paid_organization_count': sum(1 for count in org_paid.values() if count >= 2),
                'release_count': release_count, 'manually_delivered_count': manually_delivered,
                'assessments': assessments, 'unknown_costs': True, 'disclaimer': MONEY_DISCLAIMER}
