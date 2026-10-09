"""TenderOS local pilot: source-aware bid decision intelligence, not autonomous bidding.
Run with: uvicorn app.main:app --host 127.0.0.1 --port 8000
"""
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Literal, Optional

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse, Response
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get('TENDEROS_DATA_DIR', str(BASE.parent / 'data'))).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB = DATA_DIR / 'tenderos.sqlite3'
STORAGE = DATA_DIR / 'private_uploads'
STORAGE.mkdir(exist_ok=True)

app = FastAPI(title='TenderOS Local Pilot', version='0.1.0', docs_url='/api/docs', redoc_url=None)

SCHEMA = '''
PRAGMA foreign_keys=ON;
CREATE TABLE IF NOT EXISTS organizations (
 id INTEGER PRIMARY KEY, name TEXT NOT NULL, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS tenders (
 id INTEGER PRIMARY KEY, organization_id INTEGER NOT NULL REFERENCES organizations(id),
 title TEXT NOT NULL, buyer TEXT, notice_id TEXT, source_url TEXT,
 deadline TEXT, created_at TEXT NOT NULL,
 source_scope_reviewer TEXT, source_scope_note TEXT, source_scope_at TEXT, source_scope_fingerprint TEXT
);
CREATE TABLE IF NOT EXISTS sources (
 id INTEGER PRIMARY KEY, tender_id INTEGER NOT NULL REFERENCES tenders(id),
 name TEXT NOT NULL, sha256 TEXT NOT NULL, bytes INTEGER NOT NULL,
 private_path TEXT NOT NULL, pages INTEGER NOT NULL, uploaded_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS requirements (
 id INTEGER PRIMARY KEY, tender_id INTEGER NOT NULL REFERENCES tenders(id),
 source_id INTEGER REFERENCES sources(id), source_page INTEGER, source_quote TEXT,
 text TEXT NOT NULL, mandatory INTEGER, reviewed INTEGER NOT NULL DEFAULT 0,
 extraction_method TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'UNKNOWN',
 evidence_id INTEGER REFERENCES evidence(id), notes TEXT NOT NULL DEFAULT '',
 created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS evidence (
 id INTEGER PRIMARY KEY, organization_id INTEGER NOT NULL REFERENCES organizations(id),
 label TEXT NOT NULL, reference TEXT NOT NULL, expires_on TEXT,
 verified INTEGER NOT NULL DEFAULT 0, verification_note TEXT NOT NULL DEFAULT '',
 created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS decisions (
 id INTEGER PRIMARY KEY, tender_id INTEGER NOT NULL REFERENCES tenders(id),
 decision TEXT NOT NULL CHECK (decision IN ('BID','NO_BID','HOLD')),
 reviewer TEXT NOT NULL, rationale TEXT NOT NULL, decision_snapshot TEXT NOT NULL,
 created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS audit_events (
 id INTEGER PRIMARY KEY, tender_id INTEGER, actor TEXT NOT NULL,
 action TEXT NOT NULL, payload TEXT NOT NULL, created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS pilot_metrics (
 tender_id INTEGER PRIMARY KEY REFERENCES tenders(id),
 quoted_bdt REAL, paid_bdt REAL, baseline_minutes REAL, actual_minutes REAL,
 decision_changed INTEGER, repeat_customer INTEGER, benchmark_minutes REAL,
 notes TEXT NOT NULL DEFAULT '', updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_req_tender ON requirements(tender_id);
CREATE INDEX IF NOT EXISTS idx_evidence_org ON evidence(organization_id);
CREATE INDEX IF NOT EXISTS idx_audit_tender ON audit_events(tender_id);
'''


def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

@contextmanager
def conn():
    db = sqlite3.connect(DB, timeout=15)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.execute('PRAGMA busy_timeout=15000')
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

with conn() as db:
    db.executescript(SCHEMA)


def require(db, table, id):
    if table not in ('organizations', 'tenders', 'requirements', 'evidence', 'sources'):
        raise ValueError('Invalid table')
    item = db.execute(f'SELECT * FROM {table} WHERE id=?', (id,)).fetchone()
    if not item:
        raise HTTPException(404, f'{table} record not found')
    return item


def log(db, tender_id, action, payload, actor='local-operator'):
    db.execute('INSERT INTO audit_events(tender_id,actor,action,payload,created_at) VALUES(?,?,?,?,?)',
               (tender_id, actor[:100], action, json.dumps(payload, ensure_ascii=False), now()))


def asdict(row):
    return dict(row) if row else None

class OrganizationIn(BaseModel):
    name: str = Field(min_length=2, max_length=150)

class TenderIn(BaseModel):
    organization_id: int
    title: str = Field(min_length=3, max_length=300)
    buyer: str = Field(default='', max_length=250)
    notice_id: str = Field(default='', max_length=120)
    source_url: str = Field(default='', max_length=1000)
    deadline: Optional[date] = None

class RequirementIn(BaseModel):
    text: str = Field(min_length=5, max_length=5000)
    source_page: Optional[int] = Field(default=None, ge=1)
    source_quote: str = Field(default='', max_length=4000)
    mandatory: Optional[bool] = None
    reviewed: bool = False

class RequirementReview(BaseModel):
    text: str = Field(min_length=5, max_length=5000)
    mandatory: bool
    source_page: Optional[int] = Field(default=None, ge=1)
    source_quote: str = Field(default='', max_length=4000)

class EvidenceIn(BaseModel):
    organization_id: int
    label: str = Field(min_length=2, max_length=200)
    reference: str = Field(min_length=2, max_length=500)
    expires_on: Optional[date] = None
    verified: bool = False
    verification_note: str = Field(default='', max_length=2000)

class StatusIn(BaseModel):
    status: Literal['UNKNOWN', 'VERIFIED', 'PARTIAL', 'NOT_HELD']
    evidence_id: Optional[int] = None
    notes: str = Field(default='', max_length=3000)

class SourceScopeIn(BaseModel):
    reviewer: str = Field(min_length=2, max_length=100)
    checked_full_document_set: bool
    note: str = Field(min_length=15, max_length=3000)

class DecisionIn(BaseModel):
    decision: Literal['BID', 'NO_BID', 'HOLD']
    reviewer: str = Field(min_length=2, max_length=100)
    rationale: str = Field(min_length=10, max_length=3000)

class PilotMetricsIn(BaseModel):
    quoted_bdt: Optional[float] = Field(default=None, ge=0)
    paid_bdt: Optional[float] = Field(default=None, ge=0)
    baseline_minutes: Optional[float] = Field(default=None, ge=0)
    actual_minutes: Optional[float] = Field(default=None, ge=0)
    decision_changed: Optional[bool] = None
    repeat_customer: Optional[bool] = None
    benchmark_minutes: Optional[float] = Field(default=None, ge=0)
    notes: str = Field(default='', max_length=2500)


def find_evidence(db, req, evidence_id):
    if evidence_id is None:
        return None
    ev = require(db, 'evidence', evidence_id)
    ten = require(db, 'tenders', req['tender_id'])
    if ev['organization_id'] != ten['organization_id']:
        raise HTTPException(422, 'Evidence belongs to another organization')
    return ev


def evidence_current(ev):
    return bool(ev and ev['verified'] and (not ev['expires_on'] or ev['expires_on'] >= date.today().isoformat()))


def get_snapshot(db, tender_id):
    tender = asdict(require(db, 'tenders', tender_id))
    reqs = [dict(r) for r in db.execute('SELECT * FROM requirements WHERE tender_id=? ORDER BY id', (tender_id,))]
    all_verified = 0
    reviewed = 0
    counts = {'VERIFIED': 0, 'PARTIAL': 0, 'NOT_HELD': 0, 'UNKNOWN': 0}
    blockers = []
    unresolved = []
    linked_evidence = []
    for req in reqs:
        if req['reviewed'] and req['mandatory'] is not None:
            reviewed += 1
        ev = find_evidence(db, req, req['evidence_id']) if req['evidence_id'] else None
        if ev:
            linked_evidence.append({k:ev[k] for k in
                ('id','organization_id','label','reference','expires_on','verified','verification_note')})
        effective = req['status']
        if effective == 'VERIFIED' and not evidence_current(ev):
            effective = 'UNKNOWN'
            req['evidence_issue'] = 'Linked evidence is unverified or expired'
        req['effective_status'] = effective
        req['evidence_label'] = ev['label'] if ev else None
        counts[effective] += 1
        if req['mandatory'] == 1:
            if not req['source_quote'] or not (req['source_page'] or tender['notice_id'] or tender['source_url']):
                unresolved.append({'requirement_id':req['id'], 'text':req['text'], 'reason':'Source provenance incomplete'})
            if effective == 'NOT_HELD' and req['reviewed']:
                blockers.append({'requirement_id':req['id'], 'text':req['text']})
            elif not req['reviewed'] or effective != 'VERIFIED':
                unresolved.append({'requirement_id':req['id'], 'text':req['text']})
            else:
                all_verified += 1
        elif not req['reviewed'] or req['mandatory'] is None:
            unresolved.append({'requirement_id':req['id'], 'text':req['text']})
    source_ids=[dict(r) for r in db.execute('SELECT id,name,sha256,bytes,pages,uploaded_at FROM sources WHERE tender_id=? ORDER BY id',(tender_id,))]
    source_scope_fingerprint=hashlib.sha256(json.dumps({
       'sources':source_ids,
       'notice_id':tender['notice_id'], 'source_url':tender['source_url'],
       'requirements':[{k:r[k] for k in ('id','text','source_id','source_page','source_quote','mandatory','reviewed')} for r in reqs]
    },sort_keys=True,ensure_ascii=False).encode('utf-8')).hexdigest()
    source_scope_verified=bool(tender['source_scope_fingerprint'] and tender['source_scope_fingerprint']==source_scope_fingerprint)
    attestation_event=db.execute("SELECT id FROM audit_events WHERE tender_id=? AND action='human_source_scope_attested' ORDER BY id DESC LIMIT 1",(tender_id,)).fetchone()
    source_scope_attestation={'event_id':attestation_event['id'] if attestation_event else None,
        'reviewer':tender['source_scope_reviewer'], 'note':tender['source_scope_note'],
        'at':tender['source_scope_at'], 'scope_hash':tender['source_scope_fingerprint']}
    mandatory = sum(1 for r in reqs if r['mandatory'] == 1)
    fingerprint = hashlib.sha256(json.dumps({
        'requirements':[{k:r[k] for k in ('id','mandatory','reviewed','text','source_id','source_page','source_quote','status','effective_status','evidence_id','evidence_issue','notes') if k in r}
            for r in reqs],
        'linked_evidence':linked_evidence
    }, sort_keys=True, ensure_ascii=False).encode('utf-8')).hexdigest()
    if blockers:
        compliance = 'BLOCKED'
    elif not reqs or reviewed != len(reqs) or unresolved or not source_scope_verified:
        compliance = 'UNRESOLVED'
    elif all_verified == mandatory:
        compliance = 'READY_FOR_HUMAN_DECISION'
    else:
        compliance = 'UNRESOLVED'
    # Unreviewed statuses or mandatory flags always stop BID. Optional PARTIAL/UNKNOWN require explicit reviewer awareness
    # but do not by themselves disqualify when the requirement has been reviewed and classified optional.
    decisions = [dict(row) for row in db.execute('SELECT * FROM decisions WHERE tender_id=? ORDER BY id DESC', (tender_id,))]
    for i, decision in enumerate(decisions):
        try:
            past=json.loads(decision['decision_snapshot'])
        except (TypeError, ValueError):
            past={}
        past_attestation=past.get('scope_attestation') if isinstance(past,dict) else None
        attestation_matches=(isinstance(past_attestation,dict)
            and past_attestation.keys()==source_scope_attestation.keys()
            and all(type(past_attestation[k]) is type(v) and past_attestation[k]==v
                for k,v in source_scope_attestation.items()))
        # Legacy/incomplete snapshots remain historical; never rewrite stored approvals.
        decision['is_current']=(i==0 and isinstance(past,dict) and type(past.get('currency_version')) is int and past.get('currency_version')==1
            and past.get('fingerprint')==fingerprint and past.get('scope_hash')==source_scope_fingerprint
            and attestation_matches
            and (decision['decision']!='BID' or compliance=='READY_FOR_HUMAN_DECISION'))
    return {'tender': tender, 'requirements': reqs, 'counts': counts,
            'mandatory_total': mandatory, 'mandatory_verified': all_verified,
            'reviewed_count': reviewed, 'compliance': compliance,
            'blockers': blockers, 'unresolved': unresolved,
            'coverage_percent': round(100 * all_verified / mandatory, 1) if mandatory else None,
            'decisions': decisions, 'fingerprint':fingerprint,
            'current_decision': decisions[0]['decision'] if decisions and decisions[0]['is_current'] else None,
            'source_scope_verified':source_scope_verified, 'source_scope_fingerprint':source_scope_fingerprint,
            'source_scope_attestation':source_scope_attestation}

@app.middleware('http')
async def local_security_headers(request, call_next):
    # Local development only. Production must add OIDC, CSRF and tenant authorization.
    origin = request.headers.get('origin')
    if origin:
        from urllib.parse import urlparse
        host = urlparse(origin).hostname
        if host not in ('localhost', '127.0.0.1'):
            return Response(status_code=403, content='Cross-origin requests blocked')
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Content-Security-Policy'] = "default-src 'self'; img-src 'self' data:; style-src 'self'; script-src 'self'; object-src 'none'; frame-ancestors 'none'; form-action 'self'"
    response.headers['Cache-Control'] = 'no-store'
    return response

@app.get('/')
def index():
    return FileResponse(BASE / 'static' / 'index.html')

@app.get('/static/{name}')
def static(name: str):
    if name not in ('ui.js', 'style.css'):
        raise HTTPException(404)
    return FileResponse(BASE / 'static' / name)

@app.get('/api/health')
def health():
    return {'status':'ok', 'mode':'LOCAL_PILOT', 'external_model_calls':False}

@app.get('/api/organizations')
def orgs():
    with conn() as db:
        return [dict(x) for x in db.execute('SELECT * FROM organizations ORDER BY id DESC')]

@app.post('/api/organizations', status_code=201)
def create_org(item: OrganizationIn):
    with conn() as db:
        cur = db.execute('INSERT INTO organizations(name,created_at) VALUES(?,?)', (item.name.strip(), now()))
        log(db, None, 'organization_created', {'id':cur.lastrowid, 'name':item.name})
        return {'id':cur.lastrowid}

@app.get('/api/tenders')
def tenders():
    with conn() as db:
        return [dict(x) for x in db.execute('''SELECT t.*, o.name AS organization_name
            FROM tenders t JOIN organizations o ON o.id=t.organization_id ORDER BY t.id DESC''')]

@app.post('/api/tenders', status_code=201)
def create_tender(item: TenderIn):
    with conn() as db:
        require(db, 'organizations', item.organization_id)
        cur = db.execute('''INSERT INTO tenders(organization_id,title,buyer,notice_id,source_url,deadline,created_at)
             VALUES(?,?,?,?,?,?,?)''', (item.organization_id,item.title.strip(),item.buyer.strip(),
             item.notice_id.strip(),item.source_url.strip(), str(item.deadline) if item.deadline else None,now()))
        log(db, cur.lastrowid, 'tender_created', item.model_dump(mode='json'))
        return {'id':cur.lastrowid}

@app.get('/api/tenders/{tender_id}')
def tender_detail(tender_id: int):
    with conn() as db:
        out = get_snapshot(db, tender_id)
        out['sources'] = [dict(x) for x in db.execute('SELECT id,tender_id,name,sha256,bytes,pages,uploaded_at FROM sources WHERE tender_id=? ORDER BY id', (tender_id,))]
        out['evidence'] = [dict(x) for x in db.execute('SELECT * FROM evidence WHERE organization_id=? ORDER BY id DESC', (out['tender']['organization_id'],))]
        out['audit'] = [dict(x) for x in db.execute('SELECT * FROM audit_events WHERE tender_id=? ORDER BY id DESC LIMIT 50', (tender_id,))]
        out['metrics'] = asdict(db.execute('SELECT * FROM pilot_metrics WHERE tender_id=?', (tender_id,)).fetchone())
        return out

@app.post('/api/tenders/{tender_id}/requirements', status_code=201)
def create_req(tender_id: int, item: RequirementIn):
    with conn() as db:
        require(db, 'tenders', tender_id)
        stamp=now()
        cur=db.execute('''INSERT INTO requirements(tender_id,text,source_page,source_quote,mandatory,reviewed,
          extraction_method,status,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)''',
          (tender_id,item.text.strip(),item.source_page,item.source_quote.strip(),
           int(item.mandatory) if item.mandatory is not None else None,0,'manual','UNKNOWN',stamp,stamp))
        log(db,tender_id,'requirement_registered',{'id':cur.lastrowid,'page':item.source_page})
        return {'id':cur.lastrowid,'reviewed':False, 'message':'Review is required before classification is accepted'}

@app.put('/api/requirements/{req_id}/review')
def review_req(req_id: int, item: RequirementReview):
    with conn() as db:
        req=require(db,'requirements',req_id)
        db.execute('''UPDATE requirements SET text=?,mandatory=?,source_page=?,source_quote=?,reviewed=1,updated_at=? WHERE id=?''',
           (item.text.strip(),int(item.mandatory),item.source_page,item.source_quote.strip(),now(),req_id))
        log(db,req['tender_id'],'requirement_human_reviewed',{'id':req_id,'mandatory':item.mandatory, 'text':item.text[:150]})
        return {'reviewed':True}

@app.post('/api/evidence',status_code=201)
def add_evidence(item: EvidenceIn):
    with conn() as db:
        require(db,'organizations',item.organization_id)
        if item.verified and not item.verification_note.strip():
            raise HTTPException(422,'Evidence verification requires a note describing what was checked')
        cur=db.execute('''INSERT INTO evidence(organization_id,label,reference,expires_on,verified,verification_note,created_at)
         VALUES(?,?,?,?,?,?,?)''',(item.organization_id,item.label.strip(),item.reference.strip(),
         str(item.expires_on) if item.expires_on else None,int(item.verified),item.verification_note.strip(),now()))
        log(db,None,'evidence_registered',{'id':cur.lastrowid,'verified':item.verified})
        return {'id':cur.lastrowid}

@app.put('/api/requirements/{req_id}/status')
def set_status(req_id:int,item:StatusIn):
    with conn() as db:
        req=require(db,'requirements',req_id)
        if not req['reviewed'] or req['mandatory'] is None:
            raise HTTPException(422,'Human review of requirement and mandatory classification is required')
        ev=find_evidence(db,req,item.evidence_id)
        if item.status=='VERIFIED' and not evidence_current(ev):
            raise HTTPException(422,'VERIFIED requires linked, human-verified, unexpired organization evidence')
        if item.status in ('PARTIAL','NOT_HELD') and not item.notes.strip():
            raise HTTPException(422,'PARTIAL and NOT_HELD require a documented rationale')
        db.execute('UPDATE requirements SET status=?,evidence_id=?,notes=?,updated_at=? WHERE id=?',
                   (item.status,item.evidence_id,item.notes.strip(),now(),req_id))
        log(db,req['tender_id'],'requirement_status_reviewed',{'id':req_id,'status':item.status,'evidence_id':item.evidence_id})
        return {'status':item.status}

@app.post('/api/tenders/{tender_id}/attest-source-scope')
def attest_source_scope(tender_id:int,item:SourceScopeIn):
    if not item.checked_full_document_set:
        raise HTTPException(422,'Reviewer must confirm they checked all known tender documents, annexes and amendments')
    with conn() as db:
        snap=get_snapshot(db,tender_id)
        if not snap['requirements'] or any(not r['reviewed'] or r['mandatory'] is None for r in snap['requirements']):
            raise HTTPException(422,'Review every registered requirement before attesting source coverage')
        db.execute('''UPDATE tenders SET source_scope_reviewer=?,source_scope_note=?,source_scope_at=?,source_scope_fingerprint=? WHERE id=?''',
          (item.reviewer.strip(),item.note.strip(),now(),snap['source_scope_fingerprint'],tender_id))
        log(db,tender_id,'human_source_scope_attested',{'reviewer':item.reviewer.strip(),'scope_hash':snap['source_scope_fingerprint']},item.reviewer.strip())
        return {'attested':True,'scope_fingerprint':snap['source_scope_fingerprint']}

@app.post('/api/tenders/{tender_id}/decisions',status_code=201)
def decide(tender_id:int,item:DecisionIn):
    with conn() as db:
        snapshot=get_snapshot(db,tender_id)
        if item.decision=='BID' and snapshot['compliance']!='READY_FOR_HUMAN_DECISION':
            raise HTTPException(409,'BID cannot be approved: mandatory compliance is blocked or unresolved')
        decision_data={'compliance':snapshot['compliance'],'mandatory_total':snapshot['mandatory_total'],
            'mandatory_verified':snapshot['mandatory_verified'],'requirement_ids':[r['id'] for r in snapshot['requirements']],
            'fingerprint':snapshot['fingerprint'],'scope_hash':snapshot['source_scope_fingerprint'],'currency_version':1,
            'scope_attestation':snapshot['source_scope_attestation']}
        cur=db.execute('INSERT INTO decisions(tender_id,decision,reviewer,rationale,decision_snapshot,created_at) VALUES(?,?,?,?,?,?)',
           (tender_id,item.decision,item.reviewer.strip(),item.rationale.strip(),json.dumps(decision_data),now()))
        log(db,tender_id,'human_decision',{'id':cur.lastrowid,'decision':item.decision,'reviewer':item.reviewer.strip() },item.reviewer.strip())
        return {'id':cur.lastrowid,'decision':item.decision}

@app.post('/api/tenders/{tender_id}/metrics')
def save_metrics(tender_id:int,item:PilotMetricsIn):
    with conn() as db:
        require(db,'tenders',tender_id)
        db.execute('''INSERT INTO pilot_metrics(tender_id,quoted_bdt,paid_bdt,baseline_minutes,actual_minutes,
           decision_changed,repeat_customer,benchmark_minutes,notes,updated_at)
           VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(tender_id) DO UPDATE SET
           quoted_bdt=excluded.quoted_bdt,paid_bdt=excluded.paid_bdt,
           baseline_minutes=excluded.baseline_minutes,actual_minutes=excluded.actual_minutes,
           decision_changed=excluded.decision_changed,repeat_customer=excluded.repeat_customer,
           benchmark_minutes=excluded.benchmark_minutes,notes=excluded.notes,updated_at=excluded.updated_at''',
           (tender_id,item.quoted_bdt,item.paid_bdt,item.baseline_minutes,item.actual_minutes,
            int(item.decision_changed) if item.decision_changed is not None else None,
            int(item.repeat_customer) if item.repeat_customer is not None else None,
            item.benchmark_minutes,item.notes,now()))
        log(db,tender_id,'pilot_metrics_recorded',{'paid_bdt':item.paid_bdt,'actual_minutes':item.actual_minutes})
        return {'saved':True}

CANDIDATE=re.compile(r'\b(must|shall|required|mandatory|at least|minimum|eligib|qualification|similar contract|turnover|bid security|experience|key personnel|submission deadline|not less than)\b',re.I)

def extract_candidates(text, page):
    lines=[re.sub(r'\s+',' ',s).strip() for s in text.splitlines()]
    seen=set()
    result=[]
    for line in lines:
        if 22<=len(line)<=1200 and CANDIDATE.search(line) and line.lower() not in seen:
            result.append({'page':page,'quote':line})
            seen.add(line.lower())
    return result

@app.post('/api/tenders/{tender_id}/upload-pdf', status_code=201)
async def upload_pdf(tender_id:int, file:UploadFile=File(...)):
    with conn() as db:
        require(db,'tenders',tender_id)
    data=await file.read(10*1024*1024+1)
    if len(data)>10*1024*1024:
        raise HTTPException(413,'PDF limit is 10 MB')
    if not data.startswith(b'%PDF-') or not (file.filename or '').lower().endswith('.pdf'):
        raise HTTPException(415,'Only valid PDF files are allowed')
    try:
        reader=PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages)>150:
            raise HTTPException(422,'Encrypted or >150-page PDF is not supported in local pilot')
        pages=[]
        for i,page in enumerate(reader.pages,1):
            pages.append(extract_candidates(page.extract_text() or '',i))
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422,'PDF could not be safely extracted; enter requirements manually')
    sha=hashlib.sha256(data).hexdigest()
    private_name=f'{uuid.uuid4().hex}.pdf'
    path=STORAGE/private_name
    path.write_bytes(data)
    try:
        with conn() as db:
            cur=db.execute('INSERT INTO sources(tender_id,name,sha256,bytes,private_path,pages,uploaded_at) VALUES(?,?,?,?,?,?,?)',
               (tender_id,Path(file.filename).name[:180],sha,len(data),private_name,len(pages),now()))
            sid=cur.lastrowid
            stamp=now()
            count=0
            for page in pages:
                for candidate in page[:35]:
                    db.execute('''INSERT INTO requirements(tender_id,source_id,source_page,source_quote,text,
                    mandatory,reviewed,extraction_method,status,created_at,updated_at)
                    VALUES(?,?,?,?,?,NULL,0,'heuristic','UNKNOWN',?,?)''',
                    (tender_id,sid,candidate['page'],candidate['quote'],candidate['quote'],stamp,stamp))
                    count+=1
            log(db,tender_id,'pdf_uploaded_candidates_unreviewed',{'source_id':sid,'sha256':sha,'pages':len(pages),'candidates':count})
        return {'source_id':sid,'sha256':sha,'pages':len(pages),'candidates':count,'review_required':True}
    except Exception:
        path.unlink(missing_ok=True)
        raise


def report_rows(snapshot):
    return [
       {'ID':r['id'], 'Mandatory':'Yes' if r['mandatory']==1 else ('No' if r['mandatory']==0 else 'UNCLASSIFIED'),
        'Reviewed':'Yes' if r['reviewed'] else 'No','Effective status':r['effective_status'],
        'Requirement':r['text'],'Source page':r['source_page'] or '',
        'Exact source quote':r['source_quote'] or '', 'Evidence':r['evidence_label'] or '',
        'Notes':r['notes'] or ''} for r in snapshot['requirements']]

def report_guard(db,tender_id):
    snap=get_snapshot(db,tender_id)
    return snap

def spreadsheet_literal(value):
    # Excel interprets these prefixes as formulas, including after whitespace.
    return "'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')) else value

@app.get('/api/tenders/{tender_id}/export.csv')
def export_csv(tender_id:int):
    with conn() as db:
        rows=report_rows(report_guard(db,tender_id))
    out=io.StringIO()
    writer=csv.DictWriter(out,fieldnames=list(rows[0]) if rows else ['ID','Mandatory','Reviewed','Effective status','Requirement','Source page','Exact source quote','Evidence','Notes'])
    writer.writeheader()
    for row in rows:
        # Prevent spreadsheet formula injection when opening CSV in Excel.
        writer.writerow({k:spreadsheet_literal(v) for k,v in row.items()})
    return Response(out.getvalue(),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="tenderos_{tender_id}_matrix.csv"'})

@app.get('/api/tenders/{tender_id}/export.xlsx')
def export_xlsx(tender_id:int):
    with conn() as db:
        snap=report_guard(db,tender_id)
    wb=Workbook()
    ws=wb.active; ws.title='Compliance Matrix'
    ws.append(['TenderOS / Human-reviewed compliance matrix'])
    ws.append(['Tender',spreadsheet_literal(snap['tender']['title'])]);ws.append(['Compliance',snap['compliance']]);
    ws.append(['Mandatory evidence coverage',f"{snap['coverage_percent']}%" if snap['coverage_percent'] is not None else 'No mandatory requirements classified'])
    ws.append(['Note','Not a win probability. Not a substitute for reviewer sign-off.'])
    ws.append([])
    rows=report_rows(snap)
    cols=list(rows[0]) if rows else ['ID','Mandatory','Reviewed','Effective status','Requirement','Source page','Exact source quote','Evidence','Notes']
    ws.append(cols)
    for row in rows:
        # XLSX is not immune to spreadsheet formula injection: escape user-supplied literals.
        ws.append([spreadsheet_literal(row[k]) for k in cols])
    for c in ws[7]:
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='193A53')
    ws.freeze_panes='E8'
    widths={'A':11,'B':15,'C':13,'D':21,'E':75,'F':17,'G':65,'H':34,'I':40}
    for key,width in widths.items():ws.column_dimensions[key].width=width
    ws2=wb.create_sheet('Decision History')
    ws2.append(['Decision','Reviewer','Rationale','At','Snapshot'])
    for d in snap['decisions']:
        ws2.append([spreadsheet_literal(d[k]) for k in ('decision','reviewer','rationale','created_at','decision_snapshot')])
    # Excel error tokens such as #N/A must also remain literal text.
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type='s'
    buf=io.BytesIO(); wb.save(buf);buf.seek(0)
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
          headers={'Content-Disposition':f'attachment; filename="tenderos_{tender_id}.xlsx"'})

@app.get('/api/tenders/{tender_id}/export.docx')
def export_docx(tender_id:int):
    with conn() as db:
        snap=report_guard(db,tender_id)
    d=Document();d.add_heading('TenderOS | Bid Decision Intelligence',0)
    d.add_paragraph(f"Tender: {snap['tender']['title']}")
    d.add_paragraph(f"Buyer: {snap['tender']['buyer'] or 'Not provided'}")
    d.add_paragraph(f"Compliance: {snap['compliance']}")
    d.add_paragraph(f"Full source inventory attested: {'Yes' if snap['source_scope_verified'] else 'No'}")
    cover='N/A (classification incomplete)' if snap['coverage_percent'] is None else f"{snap['coverage_percent']}%"
    d.add_paragraph(f'Mandatory evidence coverage: {cover}. This is not a probability of winning.')
    d.add_heading('Requirements and evidence',level=1)
    for r in snap['requirements']:
        d.add_heading(f"R-{r['id']} | {r['effective_status']} | {'Mandatory' if r['mandatory']==1 else ('Optional' if r['mandatory']==0 else 'UNCLASSIFIED')}",level=2)
        d.add_paragraph(r['text'])
        d.add_paragraph(f"Source page: {r['source_page'] or 'not supplied'}; Quote: {r['source_quote'] or 'not supplied'}")
        d.add_paragraph(f"Reviewed: {'Yes' if r['reviewed'] else 'No'}; Evidence: {r['evidence_label'] or 'None'}; Notes: {r['notes'] or 'None'}")
    d.add_heading('Human decisions (most recent first)',level=1)
    if not snap['decisions']:d.add_paragraph('No human decision recorded. Final decision is pending.')
    for event in snap['decisions']:
        d.add_paragraph(f"{event['created_at']} | {event['decision']} | {event['reviewer']} | {event['rationale']}")
    d.add_heading('Limitations',level=1)
    d.add_paragraph('Source extractions are preliminary until reviewed. Supporting evidence is only as reliable as the human verification recorded. No external model was called. This document does not constitute an e-GP submission or legal eligibility certification.')
    buf=io.BytesIO();d.save(buf);buf.seek(0)
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers={'Content-Disposition':f'attachment; filename="tenderos_{tender_id}.docx"'})

@app.post('/api/demo', status_code=201)
def demo():
    """Synthetic example; never use as evidence of real market or compliance."""
    with conn() as db:
        org=db.execute('INSERT INTO organizations(name,created_at) VALUES(?,?)',('DEMO ONLY — Example Consulting Ltd',now())).lastrowid
        ten=db.execute('INSERT INTO tenders(organization_id,title,buyer,notice_id,created_at) VALUES(?,?,?,?,?)',
          (org,'SYNTHETIC — Technical assistance study','Illustrative Public Agency','DEMO-001',now())).lastrowid
        stamp=now()
        for text, mandatory, status, note in [
            ('At least three similar assignments completed within five years.',1,'PARTIAL','Two references currently documented; third not evidenced.'),
            ('Valid tax registration certificate must be supplied.',1,'NOT_HELD','No current certificate registered in evidence vault.'),
            ('Provide a concise organization profile.',0,'UNKNOWN','Optional, pending review')]:
            db.execute('''INSERT INTO requirements(tender_id,text,source_quote,mandatory,reviewed,extraction_method,status,notes,created_at,updated_at)
              VALUES(?,?,?,?,1,'synthetic',?,?,?,?)''',(ten,text,'Synthetic example; not an official tender',mandatory,status,note,stamp,stamp))
        log(db,ten,'synthetic_demo_loaded',{'disclaimer':'All example data are fabricated for software testing'})
        return {'organization_id':org,'tender_id':ten,'synthetic':True}
